from collections import defaultdict
from datetime import date
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.core.management.base import CommandError
from django.db import IntegrityError, connection, transaction
from django.test.utils import CaptureQueriesContext
from django.utils import timezone
from django.test import TestCase, override_settings
from openpyxl import Workbook

from .bi import obter_bi
from .models import (
    AgregadoDespesaDiario, AgregadoDespesaMensal, AuditoriaDespesa, CategoriaDespesa, ConflitoDespesa, LoteDespesa, MovimentoDespesa,
    PagamentoDespesa, SnapshotDespesaMensal, TipoDespesa, VinculoTipoDespesa,
)
from .services import (
    consolidar_lote, criar_lote, importar_arvore, importar_arvores, reconstruir_mes, salvar_vinculo,
)


def planilha(headers, linhas, *, nome="despesas.xlsx", aba="pgtdia", linha_cabecalho=4):
    wb = Workbook()
    ws = wb.active
    ws.title = aba
    for _ in range(linha_cabecalho - 1):
        ws.append([])
    ws.append(headers)
    for linha in linhas:
        ws.append(linha)
    buffer = BytesIO()
    wb.save(buffer)
    return SimpleUploadedFile(nome, buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


class DespesasTests(TestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.pasta = TemporaryDirectory()
        cls.override = override_settings(MEDIA_ROOT=cls.pasta.name)
        cls.override.enable()

    @classmethod
    def tearDownClass(cls):
        cls.override.disable()
        cls.pasta.cleanup()
        super().tearDownClass()

    def setUp(self):
        self.familia = CategoriaDespesa.objects.create(nome="DESPESAS", caminho="DESPESAS", nivel=0)
        self.folha = CategoriaDespesa.objects.create(nome="Fixas", caminho="DESPESAS\x1fFixas", nivel=1, pai=self.familia)
        self.tipo = TipoDespesa.objects.create(nome="Aluguel", chave="ALUGUEL")
        salvar_vinculo(self.tipo, self.familia, self.folha)

    def test_duas_planilhas_de_exemplo_sao_importadas_e_idempotentes(self):
        pasta = Path(__file__).resolve().parents[3]
        setores = pasta / "Arvore de Despesas.xlsx"
        importancias = pasta / "Arvore de Despesas - Copia.xlsx"
        if not setores.exists() or not importancias.exists():
            self.skipTest("Planilhas de exemplo indisponíveis")
        resultado = importar_arvores(setores, importancias)
        self.assertEqual(
            (resultado["folhas_setores"], resultado["folhas_importancias"], resultado["tipos"]),
            (40, 12, 222),
        )
        self.assertEqual(importar_arvores(setores, importancias), resultado)

    def test_duas_arvores_criam_familias_e_vinculos_idempotentes(self):
        setores = planilha(
            ["MACROGRUPO", "GRUPO GERENCIAL", "TIPO DA DESPESA"],
            [
                ["Administrativo", "Taxas", "Alvará"],
                ["Administrativo", "Taxas", "Impostos"],
            ],
            linha_cabecalho=1,
        )
        importancias = planilha(
            ["Etiqueta de Eficiência", "Tipo Financeiro", "TIPO DA DESPESA"],
            [
                ["Essencial", "Fixa", "Alvará"],
                ["Variável", "Variável", "Impostos"],
            ],
            linha_cabecalho=1,
        )

        resultado = importar_arvores(setores, importancias)

        self.assertEqual(resultado["familias"], 2)
        self.assertEqual(resultado["tipos"], 2)
        self.assertEqual(
            set(CategoriaDespesa.objects.filter(nivel=0).values_list("nome", flat=True)),
            {"DESPESAS", "SETORES", "IMPORTÂNCIAS"},
        )
        self.assertEqual(VinculoTipoDespesa.objects.filter(familia__nome="SETORES").count(), 2)
        self.assertEqual(VinculoTipoDespesa.objects.filter(familia__nome="IMPORTÂNCIAS").count(), 2)
        self.assertEqual(importar_arvores(setores, importancias), resultado)

    def test_tipo_uma_folha_por_familia_e_reclassificacao(self):
        outro = CategoriaDespesa.objects.create(nome="Variáveis", caminho="DESPESAS\x1fVariáveis", nivel=1, pai=self.familia)
        salvar_vinculo(self.tipo, self.familia, outro)
        self.assertEqual(VinculoTipoDespesa.objects.filter(tipo=self.tipo, familia=self.familia).count(), 1)
        self.assertEqual(VinculoTipoDespesa.objects.get(tipo=self.tipo).folha, outro)
        segunda = CategoriaDespesa.objects.create(nome="OPERACIONAL", caminho="OPERACIONAL", nivel=0)
        filha = CategoriaDespesa.objects.create(nome="Unidade", caminho="OPERACIONAL\x1fUnidade", nivel=1, pai=segunda)
        salvar_vinculo(self.tipo, segunda, filha)
        self.assertEqual(self.tipo.vinculos.count(), 2)

    def test_lote_sem_id_gera_identidade_e_tipo_desconhecido_bloqueia(self):
        arquivo = planilha(["EMISSAO", "DESPESA", "VALOR", "VENCIMENTO"], [[date(2026, 1, 1), "Aluguel", 100, date(2026, 1, 20)]])
        lote = criar_lote(arquivo)
        self.assertEqual(lote.prontas, 1)
        self.assertTrue(lote.linhas.first().id_origem.startswith("auto:"))
        consolidar_lote(lote)
        self.assertEqual(MovimentoDespesa.objects.count(), 1)
        arquivo = planilha(["DESPESA", "VALOR", "VENCIMENTO"], [["Tipo não cadastrado", 100, date(2026, 1, 20)]])
        lote = criar_lote(arquivo)
        self.assertEqual(lote.pendentes, 1)
        with self.assertRaises(Exception):
            consolidar_lote(lote)
        self.assertEqual(MovimentoDespesa.objects.count(), 1)

    def test_importacao_idempotente_pagamento_e_conflito_manual(self):
        headers = ["ID", "EMISSAO", "DESPESA", "FORNECEDOR", "DOC", "OBSERVACOES", "VALOR", "VENCIMENTO", "PAGTº", "VALOR PAGO", "ESP"]
        linha = ["L-1", date(2026, 1, 1), "Aluguel", "Fornecedor sem cadastro", "A", "Sala", 100, date(2026, 1, 20), date(2026, 2, 5), 40, "PIX"]
        lote = criar_lote(planilha(headers, [linha]))
        self.assertEqual(lote.prontas, 1)
        consolidar_lote(lote)
        self.assertEqual(MovimentoDespesa.objects.count(), 1)
        self.assertEqual(PagamentoDespesa.objects.count(), 1)
        self.assertEqual(AgregadoDespesaMensal.objects.get(mes=date(2026, 1, 1), base="VENCIMENTO").valor, Decimal("100"))
        self.assertEqual(AgregadoDespesaMensal.objects.get(mes=date(2026, 2, 1), base="PAGAMENTO").valor, Decimal("40"))
        novo_lote = criar_lote(planilha(headers, [linha]))
        consolidar_lote(novo_lote)
        self.assertEqual(MovimentoDespesa.objects.count(), 1)
        mov = MovimentoDespesa.objects.get()
        mov.valor = Decimal("110")
        mov.campos_manuais = ["valor"]
        mov.save()
        linha[6] = 120
        terceiro = criar_lote(planilha(headers, [linha]))
        consolidar_lote(terceiro)
        mov.refresh_from_db()
        self.assertEqual(mov.valor, Decimal("110"))
        self.assertEqual(ConflitoDespesa.objects.filter(resolvido=False).count(), 1)

    def test_snapshot_ausente_nao_vira_zero_e_folha_soma_tipos(self):
        segundo = TipoDespesa.objects.create(nome="Energia", chave="ENERGIA")
        salvar_vinculo(segundo, self.familia, self.folha)
        MovimentoDespesa.objects.create(tipo=self.tipo, vencimento=date(2026, 1, 5), valor=Decimal("100"))
        MovimentoDespesa.objects.create(tipo=segundo, vencimento=date(2026, 1, 6), valor=Decimal("25"))
        reconstruir_mes(date(2026, 1, 1), "VENCIMENTO")
        resultado = obter_bi({"familia_id": str(self.familia.pk), "visao": "mensal", "base": "VENCIMENTO", "ano": "2026"})
        folha = next(c for c in resultado["categorias"] if c["id"] == self.folha.pk)
        self.assertEqual(Decimal(folha["valores"]["2026-01-01"]), Decimal("125"))
        self.assertIsNone(folha["valores"]["2026-02-01"])
        reconstruir_mes(date(2026, 2, 1), "VENCIMENTO")
        resultado = obter_bi({"familia_id": str(self.familia.pk), "visao": "mensal", "base": "VENCIMENTO", "ano": "2026"})
        folha = next(c for c in resultado["categorias"] if c["id"] == self.folha.pk)
        self.assertEqual(Decimal(folha["valores"]["2026-02-01"]), Decimal("0"))

    def test_bi_tipos_paginados_somam_folha_e_exigem_folha_valida(self):
        tipos = [TipoDespesa(nome=f"Tipo {i:03d}", chave=f"TIPO{i:03d}") for i in range(101)]
        TipoDespesa.objects.bulk_create(tipos)
        for tipo in TipoDespesa.objects.filter(chave__startswith="TIPO"):
            salvar_vinculo(tipo, self.familia, self.folha)
            MovimentoDespesa.objects.create(tipo=tipo, vencimento=date(2025, 1, 5), valor=Decimal("1.25"))
        MovimentoDespesa.objects.create(tipo=TipoDespesa.objects.get(chave="TIPO100"), vencimento=date(2025, 1, 6), valor=Decimal("300"))
        for mes in range(1, 13):
            reconstruir_mes(date(2025, mes, 1), "VENCIMENTO")
        parametros = {"familia_id": self.familia.pk, "folha_id": self.folha.pk, "ano": 2025, "base": "VENCIMENTO"}
        primeira = self.client.get("/api/despesas/bi/tipos/", parametros)
        segunda = self.client.get("/api/despesas/bi/tipos/", {**parametros, "page": 2})
        self.assertEqual((primeira.status_code, segunda.status_code), (200, 200))
        self.assertEqual((len(primeira.json()["results"]["tipos"]), len(segunda.json()["results"]["tipos"])), (100, 2))
        self.assertEqual(primeira.json()["results"]["tipos"][0]["nome"], "Tipo 100")
        self.assertEqual([tipo["nome"] for tipo in segunda.json()["results"]["tipos"]], ["Tipo 099", "Aluguel"])
        todas = primeira.json()["results"]["tipos"] + segunda.json()["results"]["tipos"]
        soma = sum((Decimal(tipo["valores"]["2025-01-01"]) for tipo in todas), Decimal("0"))
        self.assertEqual(soma, Decimal(primeira.json()["results"]["folha"]["valores"]["2025-01-01"]))
        self.assertEqual(self.client.get("/api/despesas/bi/tipos/", {**parametros, "folha_id": self.familia.pk}).status_code, 400)

    def test_bi_tipos_visao_geral_filtro_intermediario_e_ordenacao(self):
        grupo = CategoriaDespesa.objects.create(nome="Grupo", caminho="DESPESAS\x1fGrupo", nivel=1, pai=self.familia)
        folha_a = CategoriaDespesa.objects.create(nome="Folha A", caminho="DESPESAS\x1fGrupo\x1fFolha A", nivel=2, pai=grupo)
        folha_b = CategoriaDespesa.objects.create(nome="Folha B", caminho="DESPESAS\x1fGrupo\x1fFolha B", nivel=2, pai=grupo)
        zeta = TipoDespesa.objects.create(nome="Zeta", chave="ZETA")
        alfa = TipoDespesa.objects.create(nome="Alfa", chave="ALFA")
        beta = TipoDespesa.objects.create(nome="Beta", chave="BETA")
        salvar_vinculo(zeta, self.familia, folha_a)
        salvar_vinculo(alfa, self.familia, folha_b)
        salvar_vinculo(beta, self.familia, folha_a)
        for tipo, valor in ((zeta, "150"), (alfa, "50"), (beta, "50")):
            MovimentoDespesa.objects.create(tipo=tipo, vencimento=date(2025, 1, 5), valor=Decimal(valor))
        for mes in range(1, 13):
            reconstruir_mes(date(2025, mes, 1), "VENCIMENTO")
        params = {"familia_id": self.familia.pk, "ano": 2025, "base": "VENCIMENTO"}
        geral = self.client.get("/api/despesas/bi/tipos/", params)
        self.assertEqual(geral.status_code, 200)
        self.assertEqual(geral.json()["count"], 4)
        self.assertEqual([tipo["nome"] for tipo in geral.json()["results"]["tipos"]], ["Zeta", "Alfa", "Beta", "Aluguel"])
        self.assertIsNone(geral.json()["results"]["folha"])
        grupo_resp = self.client.get("/api/despesas/bi/tipos/", {**params, "categoria_id": grupo.pk})
        self.assertEqual([tipo["nome"] for tipo in grupo_resp.json()["results"]["tipos"]], ["Zeta", "Alfa", "Beta"])
        folha_resp = self.client.get("/api/despesas/bi/tipos/", {**params, "categoria_id": folha_a.pk})
        self.assertEqual([tipo["nome"] for tipo in folha_resp.json()["results"]["tipos"]], ["Zeta", "Beta"])
        busca_resp = self.client.get("/api/despesas/bi/tipos/", {**params, "categoria_id": grupo.pk, "busca": "alf"})
        self.assertEqual([tipo["nome"] for tipo in busca_resp.json()["results"]["tipos"]], ["Alfa"])
        anual = self.client.get("/api/despesas/bi/tipos/", {**params, "visao": "anual"})
        self.assertEqual([tipo["nome"] for tipo in anual.json()["results"]["tipos"]], ["Zeta", "Alfa", "Beta", "Aluguel"])
        outra_familia = CategoriaDespesa.objects.create(nome="OUTRA", caminho="OUTRA", nivel=0)
        self.assertEqual(self.client.get("/api/despesas/bi/tipos/", {**params, "categoria_id": outra_familia.pk}).status_code, 400)
        self.assertEqual(self.client.get("/api/despesas/bi/tipos/", {**params, "folha_id": grupo.pk}).status_code, 400)
        self.assertEqual(self.client.get("/api/despesas/bi/tipos/", {**params, "categoria_id": grupo.pk}).json()["results"]["base"], "VENCIMENTO")
        self.assertEqual(len(self.client.get("/api/despesas/categorias/", {"familia_id": self.familia.pk, "search": "Folha A"}).json()), 1)

        SnapshotDespesaMensal.objects.filter(base="VENCIMENTO", mes=date(2025, 2, 1)).delete()
        indisponivel = self.client.get("/api/despesas/bi/tipos/", params).json()["results"]
        self.assertEqual(indisponivel["estado_atualizacao"], "INCOMPLETO")
        self.assertTrue(all(tipo["total"] is None for tipo in indisponivel["tipos"]))
        self.assertEqual([tipo["nome"] for tipo in indisponivel["tipos"]], ["Alfa", "Aluguel", "Beta", "Zeta"])

    def test_bi_movimentos_pagamentos_multiplos_conciliam_celula(self):
        primeiro = MovimentoDespesa.objects.create(tipo=self.tipo, vencimento=date(2026, 1, 2), valor=Decimal("100"))
        segundo = MovimentoDespesa.objects.create(tipo=self.tipo, vencimento=date(2026, 1, 3), valor=Decimal("50"))
        PagamentoDespesa.objects.create(movimento=primeiro, data=date(2026, 1, 10), valor=Decimal("10"))
        PagamentoDespesa.objects.create(movimento=primeiro, data=date(2026, 1, 20), valor=Decimal("20"))
        PagamentoDespesa.objects.create(movimento=primeiro, data=date(2026, 2, 5), valor=Decimal("70"))
        PagamentoDespesa.objects.create(movimento=segundo, data=date(2026, 1, 25), valor=Decimal("50"))
        reconstruir_mes(date(2026, 1, 1), "PAGAMENTO")
        parametros = {"familia_id": self.familia.pk, "folha_id": self.folha.pk, "tipo_id": self.tipo.pk,
                     "ano": 2026, "visao": "mensal", "base": "PAGAMENTO", "periodo": "2026-01-01"}
        resposta = self.client.get("/api/despesas/bi/movimentos/", parametros)
        self.assertEqual(resposta.status_code, 200)
        dados = resposta.json()
        self.assertTrue(dados["conciliado"])
        self.assertEqual((Decimal(dados["total_filtrado"]), Decimal(dados["total_agregado"])), (Decimal("80"), Decimal("80")))
        self.assertEqual(sum(Decimal(item["valor_periodo"]) for item in dados["results"]), Decimal("80"))
        self.assertEqual(len(next(item for item in dados["results"] if item["id"] == primeiro.pk)["pagamentos_periodo"]), 2)
        self.assertEqual(self.client.get("/api/despesas/bi/movimentos/", {**parametros, "periodo": "2026-03-01"}).status_code, 400)

    @patch("apps.despesas.bi.timezone.localdate", return_value=date(2026, 10, 7))
    def test_bi_movimentos_mensal_equivalente_usa_mesmo_corte_diario(self, _hoje):
        MovimentoDespesa.objects.create(tipo=self.tipo, vencimento=date(2026, 1, 7), valor=Decimal("10"))
        MovimentoDespesa.objects.create(tipo=self.tipo, vencimento=date(2026, 1, 8), valor=Decimal("90"))
        reconstruir_mes(date(2026, 1, 1), "VENCIMENTO")
        parametros = {"familia_id": self.familia.pk, "folha_id": self.folha.pk, "tipo_id": self.tipo.pk,
                     "ano": 2026, "visao": "mensal", "base": "VENCIMENTO", "periodo": "2026-01-01", "periodo_equivalente": 1}
        dados = self.client.get("/api/despesas/bi/movimentos/", parametros).json()
        self.assertTrue(dados["conciliado"])
        self.assertEqual((dados["count"], Decimal(dados["total_filtrado"])), (1, Decimal("10")))
        self.assertEqual(dados["periodo"]["intervalos_detalhe"], [{"inicio": "2026-01-01", "fim": "2026-01-07"}])

    @patch("apps.despesas.bi.timezone.localdate", return_value=date(2026, 10, 7))
    def test_bi_movimentos_anual_respeita_meses_prontos_sem_incluir_lacunas(self, _hoje):
        MovimentoDespesa.objects.create(tipo=self.tipo, vencimento=date(2026, 1, 4), valor=Decimal("10"))
        MovimentoDespesa.objects.create(tipo=self.tipo, vencimento=date(2026, 12, 4), valor=Decimal("30"))
        for numero in range(1, 11):
            reconstruir_mes(date(2026, numero, 1), "VENCIMENTO")
        reconstruir_mes(date(2026, 12, 1), "VENCIMENTO")
        parametros = {"familia_id": self.familia.pk, "folha_id": self.folha.pk, "tipo_id": self.tipo.pk,
                     "ano": 2026, "visao": "anual", "base": "VENCIMENTO", "periodo": "2026"}
        dados = self.client.get("/api/despesas/bi/movimentos/", parametros).json()
        self.assertTrue(dados["conciliado"])
        self.assertEqual(Decimal(dados["total_filtrado"]), Decimal("40"))
        self.assertEqual(len(dados["periodo"]["intervalos_detalhe"]), 11)

    def test_diario_concilia_por_dia_tipo_e_base_na_virada_do_mes_bissexto(self):
        segundo = TipoDespesa.objects.create(nome="Energia", chave="ENERGIA")
        primeiro = MovimentoDespesa.objects.create(tipo=self.tipo, vencimento=date(2024, 2, 29), valor=Decimal("100.25"))
        MovimentoDespesa.objects.create(tipo=self.tipo, vencimento=date(2024, 2, 29), valor=Decimal("20.10"))
        MovimentoDespesa.objects.create(tipo=segundo, vencimento=date(2024, 2, 28), valor=Decimal("30.00"))
        PagamentoDespesa.objects.create(movimento=primeiro, data=date(2024, 3, 1), valor=Decimal("60.25"))
        reconstruir_mes(date(2024, 2, 1), "VENCIMENTO")
        reconstruir_mes(date(2024, 3, 1), "PAGAMENTO")
        self.assertEqual(
            AgregadoDespesaDiario.objects.get(dia=date(2024, 2, 29), tipo=self.tipo, base="VENCIMENTO").valor,
            Decimal("120.35"),
        )
        self.assertEqual(
            AgregadoDespesaDiario.objects.get(dia=date(2024, 2, 28), tipo=segundo, base="VENCIMENTO").valor,
            Decimal("30.00"),
        )
        self.assertEqual(
            AgregadoDespesaMensal.objects.get(mes=date(2024, 2, 1), tipo=self.tipo, base="VENCIMENTO").valor,
            Decimal("120.35"),
        )
        self.assertEqual(
            AgregadoDespesaDiario.objects.get(dia=date(2024, 3, 1), tipo=self.tipo, base="PAGAMENTO").valor,
            Decimal("60.25"),
        )
        self.assertFalse(AgregadoDespesaDiario.objects.filter(dia=date(2024, 2, 29), base="PAGAMENTO").exists())
        self.assertTrue(SnapshotDespesaMensal.objects.get(mes=date(2024, 2, 1), base="VENCIMENTO").diario_atualizado_em)

    def test_carga_historica_preenche_meses_vazios_e_pode_repetir_e_conferir(self):
        mov = MovimentoDespesa.objects.create(tipo=self.tipo, vencimento=date(2024, 1, 31), valor=Decimal("100"))
        PagamentoDespesa.objects.create(movimento=mov, data=date(2024, 2, 1), valor=Decimal("75"))
        saida = StringIO()
        call_command("reconstruir_bi_despesas", inicio="2024-01", fim="2024-03", stdout=saida)
        self.assertIn("6 períodos reconstruídos", saida.getvalue())
        self.assertEqual(AgregadoDespesaDiario.objects.count(), 2)
        self.assertEqual(AgregadoDespesaMensal.objects.count(), 2)
        self.assertEqual(SnapshotDespesaMensal.objects.filter(status="PRONTO", diario_atualizado_em__isnull=False).count(), 6)
        self.assertFalse(AgregadoDespesaDiario.objects.filter(dia__month=3).exists())
        saida = StringIO()
        call_command("reconstruir_bi_despesas", inicio="2024-01", fim="2024-03", somente_sem_diario=True, stdout=saida)
        self.assertIn("0 períodos reconstruídos; 6 já tinham diário pronto", saida.getvalue())
        call_command("reconstruir_bi_despesas", inicio="2024-01", fim="2024-03", stdout=StringIO())
        self.assertEqual(AgregadoDespesaDiario.objects.count(), 2)
        self.assertEqual(SnapshotDespesaMensal.objects.count(), 6)
        saida = StringIO()
        call_command("reconstruir_bi_despesas", inicio="2024-01", fim="2024-03", conferir=True, stdout=saida)
        self.assertIn("6 períodos com diário e mensal conciliados", saida.getvalue())
        saida = StringIO()
        call_command("reconstruir_bi_despesas", conferir=True, stdout=saida)
        self.assertIn("6 períodos com diário e mensal conciliados", saida.getvalue())

    def test_snapshot_mensal_legado_sem_diario_nao_passa_na_conferencia(self):
        SnapshotDespesaMensal.objects.create(mes=date(2024, 1, 1), base="VENCIMENTO", status="PRONTO")
        with self.assertRaises(CommandError):
            call_command("reconstruir_bi_despesas", inicio="2024-01", fim="2024-01", conferir=True, stdout=StringIO())
        call_command("reconstruir_bi_despesas", inicio="2024-01", fim="2024-01", somente_sem_diario=True, stdout=StringIO())
        self.assertTrue(SnapshotDespesaMensal.objects.get(mes=date(2024, 1, 1), base="VENCIMENTO").diario_atualizado_em)

    def test_falha_na_reconstrucao_preserva_diario_e_mensal_anterior(self):
        MovimentoDespesa.objects.create(tipo=self.tipo, vencimento=date(2024, 1, 5), valor=Decimal("100"))
        reconstruir_mes(date(2024, 1, 1), "VENCIMENTO")
        anterior = SnapshotDespesaMensal.objects.get(mes=date(2024, 1, 1), base="VENCIMENTO").diario_atualizado_em
        MovimentoDespesa.objects.create(tipo=self.tipo, vencimento=date(2024, 1, 6), valor=Decimal("50"))
        with patch("apps.despesas.services.AgregadoDespesaMensal.objects.bulk_create", side_effect=RuntimeError("falha simulada")):
            with self.assertRaisesRegex(RuntimeError, "falha simulada"):
                reconstruir_mes(date(2024, 1, 1), "VENCIMENTO")
        self.assertEqual(list(AgregadoDespesaDiario.objects.values_list("dia", "valor")), [(date(2024, 1, 5), Decimal("100"))])
        self.assertEqual(AgregadoDespesaMensal.objects.get(mes=date(2024, 1, 1), base="VENCIMENTO").valor, Decimal("100"))
        snapshot = SnapshotDespesaMensal.objects.get(mes=date(2024, 1, 1), base="VENCIMENTO")
        self.assertEqual(snapshot.status, "FALHO")
        self.assertEqual(snapshot.diario_atualizado_em, anterior)
        with self.assertRaises(CommandError):
            call_command("reconstruir_bi_despesas", inicio="2024-01", fim="2024-01", conferir=True, stdout=StringIO())

    def test_chaves_unicas_e_reconstrucao_repetida_nao_duplicam_agregados(self):
        MovimentoDespesa.objects.create(tipo=self.tipo, vencimento=date(2024, 1, 5), valor=Decimal("100"))
        reconstruir_mes(date(2024, 1, 1), "VENCIMENTO")
        reconstruir_mes(date(2024, 1, 1), "VENCIMENTO")
        self.assertEqual(AgregadoDespesaDiario.objects.count(), 1)
        self.assertEqual(SnapshotDespesaMensal.objects.count(), 1)
        with self.assertRaises(IntegrityError), transaction.atomic():
            AgregadoDespesaDiario.objects.create(dia=date(2024, 1, 5), tipo=self.tipo, base="VENCIMENTO", valor=Decimal("100"))

    def test_bi_anual_nao_aceita_previa_sem_janeiro(self):
        for mes in (3, 4, 5, 6):
            reconstruir_mes(date(2025, mes, 1), "VENCIMENTO")
        resultado = obter_bi({"familia_id": str(self.familia.pk), "visao": "anual", "base": "VENCIMENTO", "ano": "2025"})
        periodo = next(item for item in resultado["periodos"] if item["periodo"] == "2025")
        self.assertEqual(periodo["estado"], "INDISPONIVEL")

    def test_bi_identifica_tipo_com_agregado_sem_vinculo_na_familia(self):
        sem_vinculo = TipoDespesa.objects.create(nome="Sem vínculo", chave="SEM VINCULO")
        MovimentoDespesa.objects.create(tipo=sem_vinculo, vencimento=date(2025, 1, 10), valor=Decimal("75"))
        reconstruir_mes(date(2025, 1, 1), "VENCIMENTO")
        resultado = obter_bi({"familia_id": str(self.familia.pk), "visao": "mensal", "base": "VENCIMENTO", "ano": "2025"})
        self.assertIn(sem_vinculo.pk, resultado["tipos_sem_vinculo"])

    def test_api_edita_importados_sem_permitir_criacao_manual(self):
        arquivo = planilha(
            ["EMISSAO", "DESPESA", "DOC", "OBSERVACOES", "VALOR", "VENCIMENTO", "PAGTº", "VALOR PAGO"],
            [[date(2026, 1, 1), "Aluguel", "X", "Sala do centro", 100, date(2026, 1, 20), date(2026, 2, 5), 40]],
        )
        consolidar_lote(criar_lote(arquivo))
        movimento_id = MovimentoDespesa.objects.get().pk
        pagamento_id = PagamentoDespesa.objects.get().pk
        resposta = self.client.post("/api/despesas/movimentos/", data={"tipo": self.tipo.pk, "vencimento": "2026-01-20", "valor": "100.00"}, content_type="application/json")
        self.assertEqual(resposta.status_code, 405)
        resposta = self.client.post("/api/despesas/pagamentos/", data={"movimento": movimento_id, "data": "2026-02-06", "valor": "10.00"}, content_type="application/json")
        self.assertEqual(resposta.status_code, 405)
        self.assertEqual(MovimentoDespesa.objects.count(), 1)
        self.assertEqual(PagamentoDespesa.objects.count(), 1)
        self.assertEqual(self.client.get("/api/despesas/movimentos/?busca=Sala%20do%20centro").json()["count"], 1)
        self.assertEqual(AgregadoDespesaDiario.objects.get(dia=date(2026, 1, 20), base="VENCIMENTO").valor, Decimal("100"))
        self.assertEqual(AgregadoDespesaDiario.objects.get(dia=date(2026, 2, 5), base="PAGAMENTO").valor, Decimal("40"))
        self.assertEqual(SnapshotDespesaMensal.objects.get(mes=date(2026, 1, 1), base="VENCIMENTO").status, "PRONTO")
        resposta = self.client.patch(f"/api/despesas/pagamentos/{pagamento_id}/", data={"valor": "100.01"}, content_type="application/json")
        self.assertEqual(resposta.status_code, 400)
        resposta = self.client.patch(f"/api/despesas/pagamentos/{pagamento_id}/", data={"valor": "60.00"}, content_type="application/json")
        self.assertEqual(resposta.status_code, 200, resposta.content)
        self.assertEqual(AgregadoDespesaDiario.objects.get(dia=date(2026, 2, 5), base="PAGAMENTO").valor, Decimal("60"))
        resposta = self.client.patch(f"/api/despesas/movimentos/{movimento_id}/", data={"vencimento": "2026-03-20", "tipo": self.tipo.pk}, content_type="application/json")
        self.assertEqual(resposta.status_code, 200, resposta.content)
        self.assertEqual(Decimal(AgregadoDespesaMensal.objects.get(mes=date(2026, 3, 1), base="VENCIMENTO").valor), Decimal("100"))
        self.assertFalse(AgregadoDespesaMensal.objects.filter(mes=date(2026, 1, 1), base="VENCIMENTO").exists())
        self.assertFalse(AgregadoDespesaDiario.objects.filter(dia=date(2026, 1, 20), base="VENCIMENTO").exists())
        self.assertEqual(AgregadoDespesaDiario.objects.get(dia=date(2026, 3, 20), base="VENCIMENTO").valor, Decimal("100"))
        self.assertEqual(self.client.get(f"/api/despesas/movimentos/{movimento_id}/").json()["saldo"], "40.00")
        self.assertEqual(self.client.delete(f"/api/despesas/pagamentos/{pagamento_id}/").status_code, 204)
        self.assertEqual(PagamentoDespesa.objects.count(), 0)
        self.assertFalse(AgregadoDespesaDiario.objects.filter(dia=date(2026, 2, 5), base="PAGAMENTO").exists())
        self.assertEqual(self.client.delete(f"/api/despesas/movimentos/{movimento_id}/").status_code, 204)
        self.assertEqual(MovimentoDespesa.objects.count(), 0)
        self.assertFalse(AgregadoDespesaDiario.objects.filter(dia=date(2026, 3, 20), base="VENCIMENTO").exists())

    def test_api_bi_e_vinculo_folha(self):
        novo = CategoriaDespesa.objects.create(nome="Outra", caminho="DESPESAS\x1fOutra", nivel=1, pai=self.familia)
        resposta = self.client.put(f"/api/despesas/tipos/{self.tipo.pk}/vinculo/", data={"familia_id": self.familia.pk, "folha_id": novo.pk}, content_type="application/json")
        self.assertEqual(resposta.status_code, 200, resposta.content)
        resposta = self.client.put(f"/api/despesas/tipos/{self.tipo.pk}/vinculo/", data={"familia_id": self.familia.pk, "folha_id": self.familia.pk}, content_type="application/json")
        self.assertEqual(resposta.status_code, 400)
        self.assertEqual(self.client.get("/api/despesas/bi/?visao=mensal&base=VENCIMENTO&ano=2026").status_code, 200)

    def test_categorias_lote_cria_atomico_e_sem_duplicar(self):
        url = "/api/despesas/categorias/lote/"
        resposta = self.client.post(url, data={"pai_id": self.familia.pk, "nomes": ["Energia", "Água"]}, content_type="application/json")
        self.assertEqual(resposta.status_code, 201, resposta.content)
        self.assertEqual(resposta.json()["total"], 2)
        self.assertEqual(CategoriaDespesa.objects.filter(pai=self.familia).count(), 3)
        self.assertEqual(AuditoriaDespesa.objects.filter(entidade="categoria", acao="CRIAR").count(), 2)
        resposta = self.client.post(url, data={"pai_id": self.familia.pk, "nomes": ["Outra", " energia "]}, content_type="application/json")
        self.assertEqual(resposta.status_code, 400)
        self.assertFalse(CategoriaDespesa.objects.filter(nome="Outra").exists())
        resposta = self.client.post(url, data={"pai_id": self.familia.pk, "nomes": ["Nova", " nova "]}, content_type="application/json")
        self.assertEqual(resposta.status_code, 400)
        self.assertFalse(CategoriaDespesa.objects.filter(nome="Nova").exists())

    def test_folha_vinculada_nao_aceita_filhas_individuais_ou_em_lote(self):
        resposta = self.client.post("/api/despesas/categorias/", data={"pai": self.folha.pk, "nome": "Filha"}, content_type="application/json")
        self.assertEqual(resposta.status_code, 400)
        resposta = self.client.post("/api/despesas/categorias/lote/", data={"pai_id": self.folha.pk, "nomes": ["Filha"]}, content_type="application/json")
        self.assertEqual(resposta.status_code, 400)
        self.assertFalse(CategoriaDespesa.objects.filter(pai=self.folha).exists())
        categorias = self.client.get("/api/despesas/categorias/").json()
        self.assertEqual(next(c for c in categorias if c["id"] == self.folha.pk)["tipos_vinculados_count"], 1)

    def test_importacao_da_arvore_nao_transforma_folha_vinculada_em_grupo(self):
        arquivo = planilha(
            ["Família", "Macrogrupo", "Grupo Gerencial", "Etiqueta de Eficiência", "Tipo Financeiro", "Categoria Raiz"],
            [["DESPESAS", "Fixas", "Novo grupo", "Etiqueta", "Folha", "Novo tipo"]],
            linha_cabecalho=1,
        )
        with self.assertRaises(Exception):
            importar_arvore(arquivo)
        self.assertFalse(CategoriaDespesa.objects.filter(pai=self.folha).exists())
        self.assertFalse(TipoDespesa.objects.filter(nome="Novo tipo").exists())

    def test_vinculos_lote_transferencia_remocao_e_bi(self):
        outra = CategoriaDespesa.objects.create(nome="Variáveis", caminho="DESPESAS\x1fVariáveis", nivel=1, pai=self.familia)
        segundo = TipoDespesa.objects.create(nome="Água", chave="ÁGUA")
        terceiro = TipoDespesa.objects.create(nome="Luz", chave="LUZ", ativo=False)
        salvar_vinculo(segundo, self.familia, self.folha)
        MovimentoDespesa.objects.create(tipo=self.tipo, vencimento=date(2026, 1, 5), valor=Decimal("100"))
        MovimentoDespesa.objects.create(tipo=segundo, vencimento=date(2026, 1, 6), valor=Decimal("25"))
        reconstruir_mes(date(2026, 1, 1), "VENCIMENTO")
        url = f"/api/despesas/categorias/{outra.pk}/vinculos/"
        dados = {"adicionar_ids": [self.tipo.pk, terceiro.pk], "remover_ids": [], "esperados": {str(self.tipo.pk): self.folha.pk, str(terceiro.pk): None}}
        resposta = self.client.post(url, data=dados, content_type="application/json")
        self.assertEqual(resposta.status_code, 200, resposta.content)
        self.assertEqual((resposta.json()["transferidos"], resposta.json()["adicionados"]), (1, 1))
        self.assertEqual(VinculoTipoDespesa.objects.filter(tipo=self.tipo, familia=self.familia).count(), 1)
        self.assertEqual(self.client.get(url).json()["vinculados_ids"], [self.tipo.pk, terceiro.pk])
        self.assertEqual(self.client.get(f"/api/despesas/tipos/?folha_id={outra.pk}").json()["count"], 2)
        bi = obter_bi({"familia_id": str(self.familia.pk), "visao": "mensal", "base": "VENCIMENTO", "ano": "2026"})
        valores = {c["id"]: Decimal(c["valores"]["2026-01-01"]) for c in bi["categorias"] if c["id"] in (self.folha.pk, outra.pk)}
        self.assertEqual(valores, {self.folha.pk: Decimal("25"), outra.pk: Decimal("100")})
        resposta = self.client.post(url, data={"adicionar_ids": [], "remover_ids": [terceiro.pk], "esperados": {str(terceiro.pk): outra.pk}}, content_type="application/json")
        self.assertEqual(resposta.status_code, 200, resposta.content)
        self.assertEqual(resposta.json()["removidos"], 1)
        self.assertFalse(VinculoTipoDespesa.objects.filter(tipo=terceiro, familia=self.familia).exists())
        self.assertEqual(AuditoriaDespesa.objects.filter(entidade="vinculo", acao="REMOVER").count(), 1)

    def test_vinculos_lote_rejeita_dados_ambiguos_sem_gravar(self):
        outra = CategoriaDespesa.objects.create(nome="Variáveis", caminho="DESPESAS\x1fVariáveis", nivel=1, pai=self.familia)
        url = f"/api/despesas/categorias/{outra.pk}/vinculos/"
        auditorias = AuditoriaDespesa.objects.count()
        casos = [
            {"adicionar_ids": [self.tipo.pk, 999999], "remover_ids": [], "esperados": {str(self.tipo.pk): self.folha.pk, "999999": None}},
            {"adicionar_ids": [self.tipo.pk], "remover_ids": [], "esperados": {str(self.tipo.pk): None}},
            {"adicionar_ids": [self.tipo.pk], "remover_ids": [self.tipo.pk], "esperados": {str(self.tipo.pk): self.folha.pk}},
        ]
        for dados in casos:
            self.assertEqual(self.client.post(url, data=dados, content_type="application/json").status_code, 400)
        self.assertEqual(VinculoTipoDespesa.objects.get(tipo=self.tipo).folha_id, self.folha.pk)
        self.assertEqual(AuditoriaDespesa.objects.count(), auditorias)

    def test_renomear_categoria_atualiza_descendentes_e_vinculos(self):
        resposta = self.client.patch(
            f"/api/despesas/categorias/{self.familia.pk}/",
            data={"nome": "DESPESAS OPERACIONAIS"}, content_type="application/json",
        )
        self.assertEqual(resposta.status_code, 200, resposta.content)
        self.folha.refresh_from_db()
        self.assertEqual(self.folha.caminho, "DESPESAS OPERACIONAIS\x1fFixas")
        self.assertEqual(VinculoTipoDespesa.objects.get(tipo=self.tipo).folha_id, self.folha.pk)
        self.assertEqual(self.client.patch(
            f"/api/despesas/categorias/{self.folha.pk}/",
            data={"pai": None}, content_type="application/json",
        ).status_code, 400)

    def test_api_lote_sem_id_consolida_com_identidade_automatica(self):
        arquivo = planilha(["EMISSAO", "DESPESA", "VALOR", "VENCIMENTO"], [[date(2026, 1, 1), "Aluguel", 100, date(2026, 1, 20)]])
        resposta = self.client.post("/api/despesas/lotes/", data={"arquivo": arquivo})
        self.assertEqual(resposta.status_code, 201, resposta.content)
        lote_id = resposta.json()["id"]
        self.assertEqual(resposta.json()["pendentes"], 0)
        self.assertEqual(LoteDespesa.objects.get(pk=lote_id).linhas.get().numero, 5)
        resposta = self.client.post(f"/api/despesas/lotes/{lote_id}/consolidar/", data={}, content_type="application/json")
        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(MovimentoDespesa.objects.count(), 1)

    def test_captura_xlsm_le_pgtdia_mesmo_com_outra_aba_ativa(self):
        wb = Workbook()
        wb.active.title = "Resumo"
        ws = wb.create_sheet("pgtdia")
        ws.append(["Relatório de pagamentos"])
        ws.append([])
        ws.append([])
        ws.append(["DESPESA", "VALOR", "VENCIMENTO"])
        ws.append(["Aluguel", 100, date(2026, 1, 20)])
        buffer = BytesIO()
        wb.save(buffer)
        arquivo = SimpleUploadedFile("nome_variavel.xlsm", buffer.getvalue())

        resposta = self.client.post("/api/despesas/lotes/", data={"arquivo": arquivo})
        self.assertEqual(resposta.status_code, 201, resposta.content)
        lote = LoteDespesa.objects.get(pk=resposta.json()["id"])
        self.assertEqual(lote.linhas.get().numero, 5)
        self.assertEqual(lote.linhas.get().dados["nome_tipo_origem"], "Aluguel")
        revalidacao = self.client.post(f"/api/despesas/lotes/{lote.pk}/revalidar/", data={}, content_type="application/json")
        self.assertEqual(revalidacao.status_code, 200, revalidacao.content)

    def test_captura_exige_aba_pgtdia_e_cabecalho_na_linha_4(self):
        arquivo = planilha(["DESPESA", "VALOR", "VENCIMENTO"], [["Aluguel", 100, date(2026, 1, 20)]], aba="PGTDIA")
        resposta = self.client.post("/api/despesas/lotes/", data={"arquivo": arquivo})
        self.assertEqual(resposta.status_code, 400)
        self.assertIn("pgtdia", str(resposta.json()))
        self.assertFalse(LoteDespesa.objects.exists())

        arquivo = planilha(["DESPESA", "VALOR", "VENCIMENTO"], [["Aluguel", 100, date(2026, 1, 20)]], linha_cabecalho=5)
        resposta = self.client.post("/api/despesas/lotes/", data={"arquivo": arquivo})
        self.assertEqual(resposta.status_code, 400)
        self.assertIn("linha 4", str(resposta.json()))
        self.assertFalse(LoteDespesa.objects.exists())

    def test_sem_id_reimporta_edicao_de_pagamento_por_chave_unica(self):
        cabecalhos = ["EMISSAO", "DESPESA", "FORNECEDOR", "DOC", "VALOR", "VENCIMENTO", "PAGTº", "VALOR PAGO"]
        original = [date(2026, 1, 1), "Aluguel", "Fornecedor", "DOC-1", 100, date(2026, 1, 20), date(2026, 1, 25), 50]
        primeiro = criar_lote(planilha(cabecalhos, [original]))
        consolidar_lote(primeiro)
        alterado = [*original]
        alterado[7] = 60
        segundo = criar_lote(planilha(cabecalhos, [alterado]))
        self.assertEqual(segundo.prontas, 1)
        self.assertEqual(segundo.linhas.first().dados["identidade_situacao"], "CONCILIADO_UNICO")
        consolidar_lote(segundo)
        self.assertEqual(MovimentoDespesa.objects.count(), 1)
        self.assertEqual(PagamentoDespesa.objects.get().valor, Decimal("60"))

    def test_identidade_ambigua_exige_revisao(self):
        cabecalhos = ["EMISSAO", "DESPESA", "FORNECEDOR", "DOC", "VALOR", "VENCIMENTO"]
        primeira = [date(2026, 1, 1), "Aluguel", "Fornecedor", "DOC-1", 100, date(2026, 1, 20)]
        segunda = [date(2026, 1, 1), "Aluguel", "Fornecedor", "DOC-1", 200, date(2026, 1, 21)]
        lote = criar_lote(planilha(cabecalhos, [primeira, segunda]))
        self.assertEqual(lote.prontas, 2)
        consolidar_lote(lote)
        nova = [*segunda]
        nova[4] = 250
        lote2 = criar_lote(planilha(cabecalhos, [primeira, nova]))
        self.assertEqual(lote2.pendentes, 0)
        self.assertEqual(lote2.linhas.get(numero=2).dados["classificacao"], "EXISTENTE")
        self.assertEqual(lote2.linhas.get(numero=3).dados["classificacao"], "ALTERADA")

    def test_ocorrencias_identicas_e_reimportacao_contam_quantidade(self):
        cabecalhos = ["EMISSAO", "DESPESA", "DOC", "VALOR", "VENCIMENTO"]
        linha = [date(2026, 1, 1), "Aluguel", "DOC-1", 100, date(2026, 1, 20)]
        consolidar_lote(criar_lote(planilha(cabecalhos, [linha])))
        segunda = criar_lote(planilha(cabecalhos, [linha, linha]))
        self.assertEqual([x.dados["classificacao"] for x in segunda.linhas.order_by("numero")], ["EXISTENTE", "NOVA"])
        self.assertEqual(segunda.prontas, 2)
        consolidar_lote(segunda)
        self.assertEqual(MovimentoDespesa.objects.count(), 2)
        terceira = criar_lote(planilha(cabecalhos, [linha, linha]))
        self.assertEqual([x.dados["classificacao"] for x in terceira.linhas.order_by("numero")], ["EXISTENTE", "EXISTENTE"])
        self.assertEqual(terceira.prontas, 2)
        resumo = self.client.get(f"/api/despesas/lotes/{terceira.pk}/resumo/").json()
        self.assertEqual(resumo["classificacoes"]["EXISTENTE"], 2)
        self.assertEqual(resumo["classificacoes"].get("NOVA", 0), 0)
        consolidar_lote(terceira)
        self.assertEqual(MovimentoDespesa.objects.count(), 2)

    def test_repeticoes_com_alteracao_ambigua_nao_viram_nova(self):
        cabecalhos = ["EMISSAO", "DESPESA", "DOC", "VALOR", "VENCIMENTO"]
        original = [date(2026, 1, 1), "Aluguel", "DOC-1", 100, date(2026, 1, 20)]
        consolidar_lote(criar_lote(planilha(cabecalhos, [original, original])))
        alterada = [*original]
        alterada[3] = 120
        captura = criar_lote(planilha(cabecalhos, [original, alterada]))
        self.assertEqual(captura.pendentes, 1)
        ambigua = captura.linhas.get(numero=3)
        self.assertEqual(ambigua.dados["classificacao"], "AMBIGUA")
        self.assertEqual(len(ambigua.dados["candidatos"]), 1)
        with self.assertRaises(Exception):
            consolidar_lote(captura)
        resposta = self.client.patch(
            f"/api/despesas/lotes/{captura.pk}/linhas/{ambigua.pk}/",
            data={"movimento_id": ambigua.dados["candidatos"][0]["movimento_id"]},
            content_type="application/json",
        )
        self.assertEqual(resposta.status_code, 200, resposta.content)
        ambigua.refresh_from_db()
        self.assertEqual(ambigua.dados["classificacao"], "ALTERADA")
        consolidar_lote(captura)
        self.assertEqual(MovimentoDespesa.objects.count(), 2)

    def test_segunda_captura_bloqueada_filtros_e_descarte(self):
        cabecalhos = ["EMISSAO", "DESPESA", "DOC", "VALOR", "VENCIMENTO"]
        linha = [date(2026, 1, 1), "Aluguel", "DOC-1", 100, date(2026, 1, 20)]
        captura = criar_lote(planilha(cabecalhos, [linha]))
        resposta = self.client.post("/api/despesas/lotes/", data={"arquivo": planilha(cabecalhos, [linha])})
        self.assertEqual(resposta.status_code, 409, resposta.content)
        self.assertEqual(len(self.client.get("/api/despesas/lotes/").json()), 1)
        self.assertEqual(self.client.get(f"/api/despesas/lotes/{captura.pk}/linhas/?classificacao=NOVA").json()["count"], 1)
        self.assertEqual(self.client.get(f"/api/despesas/lotes/{captura.pk}/linhas/?busca=DOC-1").json()["count"], 1)
        self.assertEqual(self.client.delete(f"/api/despesas/lotes/{captura.pk}/").status_code, 204)
        self.assertFalse(LoteDespesa.objects.exists())
        self.assertEqual(MovimentoDespesa.objects.count(), 0)
        self.assertEqual(self.client.post("/api/despesas/lotes/", data={"arquivo": planilha(cabecalhos, [linha])}).status_code, 201)

    def test_cabecalho_repetido_ignorado_e_data_revisada(self):
        arquivo = planilha(
            ["EMISSAO", "DESPESA", "VALOR", "VENCIMENTO", "PAGTº", "VALOR PAGO"],
            [["EMISSAO", "DESPESA", "VALOR", "VENCIMENTO", "PAGTº", "VALOR PAGO"], ["data ruim", "Aluguel", 100, date(2026, 1, 20), date(2026, 1, 25), 50]],
        )
        lote = criar_lote(arquivo)
        self.assertEqual((lote.prontas, lote.pendentes), (0, 1))
        self.assertEqual(lote.linhas.get(numero=2).status, "IGNORADO")
        linha = lote.linhas.get(numero=3)
        resposta = self.client.patch(f"/api/despesas/lotes/{lote.pk}/linhas/{linha.pk}/", data={"emissao": "2026-01-01", "tipo_id": self.tipo.pk}, content_type="application/json")
        self.assertEqual(resposta.status_code, 200, resposta.content)
        lote.refresh_from_db()
        self.assertEqual((lote.prontas, lote.pendentes), (1, 0))
        linha.refresh_from_db()
        self.assertEqual(linha.dados["classificacao"], "NOVA")
        self.assertEqual(self.client.post(f"/api/despesas/lotes/{lote.pk}/consolidar/", data={}, content_type="application/json").status_code, 200)
        self.assertEqual(MovimentoDespesa.objects.count(), 1)

    def test_captura_legada_sem_classificacao_nao_e_rotulada_invalida(self):
        arquivo = planilha(["EMISSAO", "DESPESA", "VALOR", "VENCIMENTO"], [[date(2026, 1, 1), "Aluguel", 100, date(2026, 1, 20)]])
        lote = criar_lote(arquivo)
        linha = lote.linhas.get()
        dados = dict(linha.dados)
        dados.pop("classificacao")
        linha.dados = dados
        linha.save(update_fields=["dados"])
        resumo = self.client.get(f"/api/despesas/lotes/{lote.pk}/resumo/").json()
        self.assertEqual(resumo["classificacoes"], {"SEM_CLASSIFICACAO": 1})
        resposta = self.client.get(f"/api/despesas/lotes/{lote.pk}/linhas/?classificacao=SEM_CLASSIFICACAO")
        self.assertEqual(resposta.json()["count"], 1)

    def test_cabecalho_de_outro_bloco_da_planilha_e_ignorado(self):
        arquivo = planilha(
            ["EMISSAO", "DESPESA", "VALOR", "VENCIMENTO", "PAGTº", "VALOR PAGO"],
            [["DATA C", None, "Valor Total", None, None, None]],
        )
        lote = criar_lote(arquivo)
        self.assertEqual((lote.prontas, lote.pendentes), (0, 0))
        self.assertEqual(lote.linhas.get().status, "IGNORADO")

    def test_valor_pago_invalido_nao_vira_zero_e_emissao_requer_correcao(self):
        arquivo = planilha(
            ["EMISSAO", "DESPESA", "VALOR", "VENCIMENTO", "VALOR PAGO"],
            [["data ruim", "Aluguel", 100, date(2026, 1, 20), "valor ruim"]],
        )
        lote = criar_lote(arquivo)
        linha = lote.linhas.get()
        self.assertIn("Valor pago inválido", linha.erros)
        resposta = self.client.patch(
            f"/api/despesas/lotes/{lote.pk}/linhas/{linha.pk}/",
            data={"emissao": None, "valor_pago": "0"}, content_type="application/json",
        )
        self.assertEqual(resposta.status_code, 400)

    def test_revalidacao_atualiza_lote_sem_criar_copia(self):
        arquivo = planilha(
            ["EMISSAO", "DESPESA", "VALOR", "VENCIMENTO", "PAGTº", "VALOR PAGO"],
            [[date(2026, 1, 1), "Aluguel", 100, date(2026, 1, 20), date(2026, 1, 25), 50]],
        )
        lote = criar_lote(arquivo)
        linha = lote.linhas.get()
        linha.status = "PENDENTE"
        linha.erros = ["ID estável ausente"]
        linha.save(update_fields=["status", "erros"])
        lote.prontas, lote.pendentes = 0, 1
        lote.save(update_fields=["prontas", "pendentes"])
        resposta = self.client.post(f"/api/despesas/lotes/{lote.pk}/revalidar/", data={}, content_type="application/json")
        self.assertEqual(resposta.status_code, 200, resposta.content)
        self.assertEqual(resposta.json()["id"], lote.pk)
        self.assertEqual(resposta.json()["prontas"], 1)
        self.assertEqual(LoteDespesa.objects.count(), 1)

    def test_consolidacao_repetida_retorna_estado_sem_reprocessar(self):
        arquivo = planilha(
            ["EMISSAO", "DESPESA", "VALOR", "VENCIMENTO", "PAGTº", "VALOR PAGO"],
            [[date(2026, 1, 1), "Aluguel", 100, date(2026, 1, 20), date(2026, 1, 25), 50]],
        )
        lote = criar_lote(arquivo)
        consolidar_lote(lote)
        auditorias = AuditoriaDespesa.objects.count()
        resposta = consolidar_lote(lote)
        self.assertTrue(resposta["ja_consolidado"])
        self.assertEqual(AuditoriaDespesa.objects.count(), auditorias)
        self.assertEqual(MovimentoDespesa.objects.count(), 1)

    def test_reimportacao_igual_nao_recria_agregados(self):
        cabecalhos = ["EMISSAO", "DESPESA", "VALOR", "VENCIMENTO", "PAGTº", "VALOR PAGO"]
        linha = [date(2026, 1, 1), "Aluguel", 100, date(2026, 1, 20), date(2026, 1, 25), 50]
        consolidar_lote(criar_lote(planilha(cabecalhos, [linha])))
        lote = criar_lote(planilha(cabecalhos, [linha]))
        auditorias = AuditoriaDespesa.objects.count()
        resposta = consolidar_lote(lote)
        self.assertEqual(resposta["periodos"], 0)
        self.assertEqual(AuditoriaDespesa.objects.count(), auditorias)
        self.assertEqual(MovimentoDespesa.objects.count(), 1)


class DespesasBIContratoTests(TestCase):
    def setUp(self):
        self.familia = CategoriaDespesa.objects.create(nome="DESPESAS", caminho="DESPESAS", nivel=0)
        self.folha = CategoriaDespesa.objects.create(nome="Fixas", caminho="DESPESAS\x1fFixas", nivel=1, pai=self.familia)
        self.tipo = TipoDespesa.objects.create(nome="Aluguel", chave="ALUGUEL")
        salvar_vinculo(self.tipo, self.familia, self.folha)

    def pronto(self, ano, mes, base="VENCIMENTO", lancamentos=(), diario=True):
        instante = timezone.now()
        SnapshotDespesaMensal.objects.create(
            mes=date(ano, mes, 1), base=base, status="PRONTO",
            atualizado_em=instante, diario_atualizado_em=instante if diario else None,
        )
        por_dia = defaultdict(Decimal)
        por_tipo = defaultdict(Decimal)
        for dia, valor, tipo in lancamentos:
            por_dia[(dia, tipo.pk)] += Decimal(str(valor))
            por_tipo[tipo.pk] += Decimal(str(valor))
        AgregadoDespesaDiario.objects.bulk_create([
            AgregadoDespesaDiario(dia=dia, tipo_id=tipo_id, base=base, valor=valor)
            for (dia, tipo_id), valor in por_dia.items()
        ])
        AgregadoDespesaMensal.objects.bulk_create([
            AgregadoDespesaMensal(mes=date(ano, mes, 1), tipo_id=tipo_id, base=base, valor=valor)
            for tipo_id, valor in por_tipo.items()
        ])

    def bi(self, *, visao="mensal", base="VENCIMENTO", ano=2026, equivalente=False):
        return obter_bi({
            "familia_id": str(self.familia.pk), "visao": visao, "base": base,
            "ano": str(ano), "periodo_equivalente": "1" if equivalente else "0",
        })

    def folha_bi(self, dados):
        return next(item for item in dados["categorias"] if item["id"] == self.folha.pk)

    @patch("apps.despesas.bi.timezone.localdate", return_value=date(2026, 10, 7))
    def test_mensal_equivalente_corta_dias_e_mantem_vencimentos_futuros(self, _hoje):
        eventos = {
            1: [(date(2026, 1, 7), 100, self.tipo), (date(2026, 1, 20), 200, self.tipo)],
            2: [(date(2026, 2, 7), 50, self.tipo), (date(2026, 2, 28), 80, self.tipo)],
            10: [(date(2026, 10, 7), 40, self.tipo), (date(2026, 10, 8), 60, self.tipo)],
            11: [(date(2026, 11, 1), 70, self.tipo)],
        }
        for mes in range(1, 13):
            self.pronto(2026, mes, lancamentos=eventos.get(mes, ()))
        equivalente = self.bi(equivalente=True)
        valores = self.folha_bi(equivalente)
        self.assertEqual(valores["valores"]["2026-01-01"], "100.00")
        self.assertEqual(valores["valores"]["2026-02-01"], "50.00")
        self.assertEqual(valores["valores"]["2026-10-01"], "40.00")
        self.assertEqual(valores["valores"]["2026-11-01"], "70.00")
        self.assertEqual(valores["total"], "190.00")
        self.assertEqual(valores["programado_apos_corte"], "130.00")
        self.assertEqual(equivalente["periodos"][9]["data_corte"], "2026-10-07")
        self.assertTrue(equivalente["periodos"][10]["programado"])
        completo = self.bi()
        self.assertEqual(self.folha_bi(completo)["valores"]["2026-01-01"], "300.00")
        self.assertEqual(self.folha_bi(completo)["total"], "600.00")
        self.assertFalse(completo["periodo_equivalente"])

    @patch("apps.despesas.bi.timezone.localdate", return_value=date(2026, 10, 7))
    def test_mensal_historico_ignora_equivalencia_e_diario_ausente_nao_vira_zero(self, _hoje):
        self.pronto(2025, 1, lancamentos=[(date(2025, 1, 20), 100, self.tipo)], diario=False)
        historico = self.bi(ano=2025, equivalente=True)
        self.assertFalse(historico["periodo_equivalente"])
        self.assertEqual(self.folha_bi(historico)["valores"]["2025-01-01"], "100.00")
        self.pronto(2026, 1, lancamentos=[(date(2026, 1, 20), 100, self.tipo)], diario=False)
        cortado = self.bi(equivalente=True)
        self.assertIsNone(self.folha_bi(cortado)["valores"]["2026-01-01"])
        self.assertEqual(cortado["periodos"][0]["estado_snapshot"], "DIARIO_AUSENTE")
        self.assertIsNone(self.folha_bi(cortado)["total"])

    @patch("apps.despesas.bi.timezone.localdate", return_value=date(2026, 10, 7))
    def test_pagamento_futuro_fica_fora_do_fluxo_real_mesmo_sem_equivalencia(self, _hoje):
        self.pronto(2026, 10, base="PAGAMENTO", lancamentos=[
            (date(2026, 10, 6), 20, self.tipo), (date(2026, 10, 8), 30, self.tipo),
        ])
        self.pronto(2026, 11, base="PAGAMENTO", lancamentos=[(date(2026, 11, 1), 50, self.tipo)])
        dados = self.bi(base="PAGAMENTO")
        folha = self.folha_bi(dados)
        self.assertEqual(folha["valores"]["2026-10-01"], "20.00")
        self.assertIsNone(folha["valores"]["2026-11-01"])
        self.assertEqual(folha["pagamentos_futuros_registrados"], "80.00")
        self.assertEqual(dados["periodos"][10]["estado"], "FUTURO")
        SnapshotDespesaMensal.objects.filter(mes=date(2026, 10, 1), base="PAGAMENTO").update(diario_atualizado_em=None)
        sem_diario = self.bi(base="PAGAMENTO")
        self.assertIsNone(self.folha_bi(sem_diario)["valores"]["2026-10-01"])
        self.assertEqual(sem_diario["periodos"][9]["estado_snapshot"], "DIARIO_AUSENTE")

    @patch("apps.despesas.bi.timezone.localdate", return_value=date(2026, 10, 7))
    def test_anual_pagamento_sem_equivalencia_combina_meses_e_corte_atual(self, _hoje):
        for ano in (2025, 2026):
            for mes in range(1, 13):
                eventos = ()
                if mes == 1:
                    eventos = [(date(ano, 1, 5), 100, self.tipo)]
                if ano == 2026 and mes == 10:
                    eventos = [
                        (date(2026, 10, 6), 20, self.tipo),
                        (date(2026, 10, 8), 30, self.tipo),
                    ]
                if ano == 2026 and mes == 11:
                    eventos = [(date(2026, 11, 1), 50, self.tipo)]
                self.pronto(ano, mes, base="PAGAMENTO", lancamentos=eventos)
        dados = self.bi(visao="anual", base="PAGAMENTO")
        folha = self.folha_bi(dados)
        self.assertEqual(folha["valores"], {"2025": "100.00", "2026": "120.00"})
        self.assertEqual(folha["variacao_percentual"], "20.00")
        self.assertTrue(folha["variacao_recente_neutra"])
        self.assertEqual(folha["pagamentos_futuros_registrados"], "80.00")
        self.assertEqual(dados["periodos"][-1]["fim_efetivo"], "2026-10-07")
        self.assertEqual(dados["periodos"][-1]["meses_indisponiveis"], [])
        SnapshotDespesaMensal.objects.filter(mes=date(2026, 10, 1), base="PAGAMENTO").update(diario_atualizado_em=None)
        indisponivel = self.bi(visao="anual", base="PAGAMENTO")
        self.assertIsNone(self.folha_bi(indisponivel)["valores"]["2026"])
        self.assertIsNone(self.folha_bi(indisponivel)["variacao_percentual"])

    @patch("apps.despesas.bi.timezone.localdate", return_value=date(2026, 10, 7))
    def test_anual_taxas_quatro_tres_dois_e_um_ano(self, _hoje):
        for ano, valor in ((2022, 100), (2023, 120), (2024, 180), (2025, 270)):
            for mes in range(1, 13):
                lancamentos = [(date(ano, 1, 5), valor, self.tipo)] if mes == 1 else ()
                self.pronto(ano, mes, lancamentos=lancamentos)
        quatro = self.bi(visao="anual", ano=2025)
        self.assertEqual([p["periodo"] for p in quatro["periodos"]], ["2022", "2023", "2024", "2025"])
        self.assertEqual(self.folha_bi(quatro)["variacao_historica_percentual"], "35.00")
        self.assertEqual(self.folha_bi(quatro)["variacao_percentual"], "50.00")
        tres = self.bi(visao="anual", ano=2024)
        self.assertEqual(self.folha_bi(tres)["variacao_historica_percentual"], "20.00")
        dois = self.bi(visao="anual", ano=2023)
        self.assertIsNone(self.folha_bi(dois)["variacao_historica_percentual"])
        self.assertEqual(self.folha_bi(dois)["variacao_percentual"], "20.00")
        um = self.bi(visao="anual", ano=2022)
        self.assertIsNone(self.folha_bi(um)["variacao_historica_percentual"])
        self.assertIsNone(self.folha_bi(um)["variacao_percentual"])

    @patch("apps.despesas.bi.timezone.localdate", return_value=date(2026, 10, 7))
    def test_anual_base_zero_invalida_taxa_historica_sem_descartar_transicao(self, _hoje):
        for ano, valor in ((2023, 0), (2024, 100), (2025, 120)):
            for mes in range(1, 13):
                self.pronto(ano, mes, lancamentos=[(date(ano, 1, 5), valor, self.tipo)] if mes == 1 and valor else ())
        folha = self.folha_bi(self.bi(visao="anual", ano=2025))
        self.assertIsNone(folha["variacao_historica_percentual"])
        self.assertEqual(folha["variacao_percentual"], "20.00")

    @patch("apps.despesas.bi.timezone.localdate", return_value=date(2024, 2, 29))
    def test_anual_equivalente_ajusta_fevereiro_bissexto(self, _hoje):
        for ano in (2023, 2024):
            for mes in (1, 2):
                evento = [(date(ano, 2, 28 if ano == 2023 else 29), 100 if ano == 2023 else 110, self.tipo)] if mes == 2 else ()
                self.pronto(ano, mes, lancamentos=evento)
        dados = self.bi(visao="anual", ano=2024, equivalente=True)
        self.assertEqual([p["data_corte"] for p in dados["periodos"]], ["2023-02-28", "2024-02-29"])
        self.assertEqual(self.folha_bi(dados)["variacao_percentual"], "10.00")
        self.assertEqual(self.folha_bi(dados)["valores"], {"2023": "100.00", "2024": "110.00"})
        sem_corte = self.bi(visao="anual", ano=2024)
        self.assertIsNone(self.folha_bi(sem_corte)["valores"]["2023"])

    @patch("apps.despesas.bi.timezone.localdate", return_value=date(2026, 10, 7))
    def test_anual_atual_compara_corte_e_identifica_taxa_sem_corte_como_neutra(self, _hoje):
        for ano in (2025, 2026):
            for mes in range(1, 13):
                eventos = ()
                if mes == 10:
                    eventos = [(date(ano, 10, 7), 100 if ano == 2025 else 120, self.tipo),
                               (date(ano, 10, 8), 100 if ano == 2025 else 60, self.tipo)]
                if ano == 2026 and mes == 11:
                    eventos = [(date(2026, 11, 1), 70, self.tipo)]
                self.pronto(ano, mes, lancamentos=eventos)
        equivalente = self.bi(visao="anual", equivalente=True)
        self.assertEqual(self.folha_bi(equivalente)["valores"], {"2025": "100.00", "2026": "120.00"})
        self.assertEqual(self.folha_bi(equivalente)["variacao_percentual"], "20.00")
        self.assertEqual(self.folha_bi(equivalente)["programado_apos_corte"], "130.00")
        completo = self.bi(visao="anual")
        self.assertEqual(self.folha_bi(completo)["valores"], {"2025": "200.00", "2026": "250.00"})
        self.assertEqual(self.folha_bi(completo)["variacao_percentual"], "25.00")
        self.assertTrue(self.folha_bi(completo)["variacao_recente_neutra"])
        self.assertEqual(completo["periodos"][-1]["estado"], "PARCIAL")

    @patch("apps.despesas.bi.timezone.localdate", return_value=date(2026, 10, 7))
    def test_ano_futuro_mostra_somente_programado_conhecido_sem_taxas(self, _hoje):
        self.pronto(2027, 1, lancamentos=[(date(2027, 1, 5), 100, self.tipo)])
        dados = self.bi(visao="anual", ano=2027)
        folha = self.folha_bi(dados)
        self.assertEqual(dados["periodos"][-1]["posicao_calendario"], "FUTURO")
        self.assertEqual(dados["periodos"][-1]["estado"], "PARCIAL")
        self.assertEqual(folha["valores"]["2027"], "100.00")
        self.assertEqual(folha["programado_apos_corte"], "100.00")
        self.assertIsNone(folha["variacao_historica_percentual"])
        self.assertIsNone(folha["variacao_percentual"])

    @patch("apps.despesas.bi.timezone.localdate", return_value=date(2026, 10, 7))
    def test_falha_snapshot_zero_pronto_e_ausencia_sao_distintos(self, _hoje):
        self.pronto(2025, 1, lancamentos=[(date(2025, 1, 5), 100, self.tipo)])
        SnapshotDespesaMensal.objects.filter(mes=date(2025, 1, 1)).update(status="FALHO")
        self.pronto(2025, 2)
        dados = self.bi(ano=2025)
        folha = self.folha_bi(dados)
        self.assertIsNone(folha["valores"]["2025-01-01"])
        self.assertEqual(folha["valores"]["2025-02-01"], "0")
        self.assertIsNone(folha["valores"]["2025-03-01"])
        self.assertEqual([p["estado"] for p in dados["periodos"][:3]], ["FALHO", "PRONTO", "AUSENTE"])
        self.assertEqual(dados["estado_atualizacao"], "DESATUALIZADO")

    @patch("apps.despesas.bi.timezone.localdate", return_value=date(2026, 10, 7))
    def test_tipos_sem_vinculo_ou_vinculo_invalido_nao_somem_do_total(self, _hoje):
        sem_vinculo = TipoDespesa.objects.create(nome="Outro", chave="OUTRO")
        CategoriaDespesa.objects.create(nome="Subgrupo", caminho="DESPESAS\x1fFixas\x1fSubgrupo", nivel=2, pai=self.folha)
        self.pronto(2025, 1, lancamentos=[
            (date(2025, 1, 5), 100, self.tipo), (date(2025, 1, 6), 50, sem_vinculo),
        ])
        dados = self.bi(ano=2025)
        self.assertEqual(self.folha_bi(dados)["valores"]["2025-01-01"], "0")
        self.assertEqual(dados["nao_classificados"]["valores"]["2025-01-01"], "150.00")
        self.assertEqual(dados["tipos_sem_vinculo"], sorted([self.tipo.pk, sem_vinculo.pk]))
        self.assertEqual(dados["nao_classificados"]["tipos_com_vinculo_invalido"], [self.tipo.pk])

    @patch("apps.despesas.bi.timezone.localdate", return_value=date(2026, 10, 7))
    def test_api_preserva_contrato_e_nao_le_lancamentos_brutos(self, _hoje):
        self.pronto(2025, 1, lancamentos=[(date(2025, 1, 5), 100, self.tipo)])
        with CaptureQueriesContext(connection) as consultas:
            resposta = self.client.get(f"/api/despesas/bi/?familia_id={self.familia.pk}&visao=mensal&base=VENCIMENTO&ano=2025")
        self.assertEqual(resposta.status_code, 200, resposta.content)
        self.assertTrue({"familia", "visao", "base", "ano", "anos_disponiveis", "periodos", "categorias", "tipos", "tipos_sem_vinculo"} <= resposta.json().keys())
        sql = " ".join(item["sql"].lower() for item in consultas)
        self.assertNotIn("despesas_movimentodespesa", sql)
        self.assertNotIn("despesas_pagamentodespesa", sql)
        self.assertEqual(self.client.get("/api/despesas/bi/?periodo_equivalente=2").status_code, 400)
