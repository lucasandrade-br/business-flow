"""Regras de importação e consolidação das despesas.

Todas as operações usam o banco da filial selecionado pelo processo Django.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
import json
import unicodedata
from pathlib import Path

from django.core.exceptions import ValidationError
from django.core.files.base import ContentFile
from django.db import transaction
from django.db import DatabaseError, connection
from django.db.models import Max, Sum
from django.utils import timezone
from openpyxl import load_workbook

from apps.cadastros.models import Fornecedor
from .models import (
    AgregadoDespesaDiario, AgregadoDespesaMensal, AuditoriaDespesa, CategoriaDespesa, ConflitoDespesa,
    LinhaDespesa, LoteDespesa, MovimentoDespesa, PagamentoDespesa,
    SnapshotDespesaMensal, TipoDespesa, VinculoTipoDespesa,
)


def chave(texto):
    return " ".join(str(texto or "").strip().upper().split())


def cabecalho(texto):
    original = unicodedata.normalize("NFKD", str(texto or "").upper())
    return "".join(c for c in original if c.isalnum()).upper()


def mes_inicio(data):
    return date(data.year, data.month, 1)


def proximo_mes(mes):
    return date(mes.year + (mes.month == 12), mes.month % 12 + 1, 1)


def auditar(entidade, objeto_id, acao, antes=None, depois=None, operador=""):
    AuditoriaDespesa.objects.create(
        entidade=entidade, entidade_id=objeto_id, acao=acao,
        antes=antes or {}, depois=depois or {}, operador=operador or "",
    )


def validar_folha(folha: CategoriaDespesa, familia: CategoriaDespesa):
    vinculo = VinculoTipoDespesa(tipo=TipoDespesa(), familia=familia, folha=folha)
    vinculo.clean()


def verificar_pai_sem_vinculos(pai):
    if pai and VinculoTipoDespesa.objects.filter(folha=pai).exists():
        raise ValidationError({"pai": "Esta categoria possui tipos vinculados. Reclassifique ou remova os vínculos antes de criar filhas."})


def salvar_vinculo(tipo, familia, folha, operador=""):
    with transaction.atomic():
        # A criação de filhas bloqueia o mesmo nó antes de verificar vínculos.
        folha = CategoriaDespesa.objects.select_for_update().get(pk=folha.pk)
        validar_folha(folha, familia)
        TipoDespesa.objects.select_for_update().get(pk=tipo.pk)
        anterior = VinculoTipoDespesa.objects.select_for_update().filter(tipo=tipo, familia=familia).first()
        if anterior and anterior.folha_id == folha.pk:
            return anterior
        antes = {"folha_id": anterior.folha_id} if anterior else {}
        vinculo, _ = VinculoTipoDespesa.objects.update_or_create(
            tipo=tipo, familia=familia, defaults={"folha": folha},
        )
        auditar("vinculo", vinculo.pk, "CLASSIFICAR", antes, {"folha_id": folha.pk}, operador)
        return vinculo


def _ler_planilha_arvore(arquivo, colunas, descricao):
    workbook = load_workbook(arquivo, read_only=True, data_only=True)
    linhas = list(workbook.active.values)
    if not linhas:
        raise ValidationError(f"Planilha de {descricao} vazia.")
    mapa = {cabecalho(nome): i for i, nome in enumerate(linhas[0])}
    colunas_normalizadas = [cabecalho(coluna) for coluna in colunas]
    ausentes = [coluna for coluna in colunas_normalizadas if coluna not in mapa]
    if ausentes:
        nomes = ", ".join(ausentes)
        raise ValidationError(f"Planilha de {descricao} sem as colunas obrigatórias: {nomes}.")

    itens = []
    tipos_vistos = {}
    for numero, linha in enumerate(linhas[1:], 2):
        if not any(valor not in (None, "") for valor in linha):
            continue
        valores = [str(linha[mapa[coluna]] or "").strip() for coluna in colunas_normalizadas]
        if not all(valores):
            raise ValidationError(f"Planilha de {descricao}, linha {numero}: caminho ou tipo incompleto.")
        caminho = tuple(valores[:-1])
        nome_tipo = valores[-1]
        chave_tipo = chave(nome_tipo)
        if chave_tipo in tipos_vistos and tipos_vistos[chave_tipo] != caminho:
            raise ValidationError(
                f"Planilha de {descricao}, linha {numero}: o tipo está em duas folhas da mesma família."
            )
        tipos_vistos[chave_tipo] = caminho
        itens.append((caminho, nome_tipo))
    return itens


def _importar_familia(itens, nome_familia, operador):
    cache = {}
    folhas = set()
    tipos = set()
    familia, _ = CategoriaDespesa.objects.get_or_create(
        caminho=nome_familia,
        defaults={"nome": nome_familia, "nivel": 0},
    )
    if familia.nome != nome_familia or familia.nivel != 0 or familia.pai_id is not None:
        raise ValidationError(f"A família existente {nome_familia!r} é incompatível.")

    for caminho, nome_tipo in itens:
        caminho_completo = (nome_familia,) + caminho
        pai = familia
        for nivel, nome in enumerate(caminho_completo[1:], 1):
            subcaminho = "\x1f".join(caminho_completo[:nivel + 1])
            if len(subcaminho) > 255:
                raise ValidationError("Caminho do plano de despesas excede 255 caracteres.")
            if subcaminho not in cache:
                if pai and not CategoriaDespesa.objects.filter(caminho=subcaminho).exists():
                    pai = CategoriaDespesa.objects.select_for_update().get(pk=pai.pk)
                    verificar_pai_sem_vinculos(pai)
                no, _ = CategoriaDespesa.objects.get_or_create(
                    caminho=subcaminho,
                    defaults={"nome": nome, "nivel": nivel, "pai": pai},
                )
                if no.nome != nome or no.nivel != nivel or no.pai_id != pai.pk:
                    raise ValidationError(f"Árvore existente incompatível em {subcaminho!r}.")
                cache[subcaminho] = no
            pai = cache[subcaminho]

        folhas.add(pai.pk)
        tipo, criado = TipoDespesa.objects.get_or_create(
            chave=chave(nome_tipo),
            defaults={"nome": nome_tipo},
        )
        if criado:
            auditar("tipo", tipo.pk, "CRIAR", depois={"nome": nome_tipo}, operador=operador)
        tipos.add(tipo.pk)
        salvar_vinculo(tipo, familia, pai, operador)
    return familia, folhas, tipos


def importar_arvores(arquivo_setores, arquivo_importancias, operador="importacao"):
    """Importa as famílias SETORES e IMPORTÂNCIAS a partir de duas planilhas."""
    setores = _ler_planilha_arvore(
        arquivo_setores,
        ["MACROGRUPO", "GRUPOGERENCIAL", "TIPODADESPESA"],
        "setores",
    )
    importancias = _ler_planilha_arvore(
        arquivo_importancias,
        ["ETIQUETADEEFICIENCIA", "TIPOFINANCEIRO", "TIPODADESPESA"],
        "importâncias",
    )
    tipos_setores = {chave(item[1]) for item in setores}
    tipos_importancias = {chave(item[1]) for item in importancias}
    if tipos_setores != tipos_importancias:
        somente_setores = sorted(tipos_setores - tipos_importancias)
        somente_importancias = sorted(tipos_importancias - tipos_setores)
        raise ValidationError(
            "As planilhas precisam ter os mesmos tipos de despesa. "
            f"Apenas em setores: {somente_setores}; apenas em importâncias: {somente_importancias}."
        )

    with transaction.atomic():
        familia_setores, folhas_setores, tipos_setores = _importar_familia(
            setores, "SETORES", operador,
        )
        familia_importancias, folhas_importancias, tipos_importancias = _importar_familia(
            importancias, "IMPORTÂNCIAS", operador,
        )
    return {
        "familias": 2,
        "folhas": len(folhas_setores | folhas_importancias),
        "folhas_setores": len(folhas_setores),
        "folhas_importancias": len(folhas_importancias),
        "tipos": len(tipos_setores | tipos_importancias),
        "linhas": len(setores) + len(importancias),
        "vinculos": len(tipos_setores) + len(tipos_importancias),
    }


def importar_arvore(arquivo, operador="importacao"):
    """Carga idempotente da árvore; rejeita ambiguidade antes de gravar."""
    workbook = load_workbook(arquivo, read_only=True, data_only=True)
    linhas = list(workbook.active.values)
    if not linhas:
        raise ValidationError("Planilha de árvore vazia.")
    esperadas = ["FAMILIA", "MACROGRUPO", "GRUPOGERENCIAL", "ETIQUETADEEFICIENCIA", "TIPOFINANCEIRO", "CATEGORIARAIZ"]
    mapa = {cabecalho(nome): i for i, nome in enumerate(linhas[0])}
    if not all(nome in mapa for nome in esperadas):
        raise ValidationError("A árvore precisa das colunas Família, Macrogrupo, Grupo Gerencial, Etiqueta de Eficiência, Tipo Financeiro e Categoria Raiz.")
    itens = []
    tipo_para_familia = {}
    for numero, linha in enumerate(linhas[1:], 2):
        if not any(valor not in (None, "") for valor in linha):
            continue
        nomes = [str(linha[mapa[coluna]] or "").strip() for coluna in esperadas]
        if not all(nomes):
            raise ValidationError(f"Linha {numero}: caminho ou tipo incompleto.")
        caminho = tuple(nomes[:5])
        par = (chave(nomes[-1]), chave(nomes[0]))
        if par in tipo_para_familia and tipo_para_familia[par] != caminho:
            raise ValidationError(f"Linha {numero}: o tipo está em duas folhas da mesma família.")
        tipo_para_familia[par] = caminho
        codigo = str(linha[mapa["COD"]] or "").strip() if "COD" in mapa else ""
        itens.append((caminho, nomes[-1], codigo))
    with transaction.atomic():
        cache = {}
        folhas = set()
        tipos = set()
        for caminho, nome_tipo, codigo in itens:
            pai = None
            for nivel, nome in enumerate(caminho):
                subcaminho = "\x1f".join(caminho[:nivel + 1])
                if len(subcaminho) > 255:
                    raise ValidationError("Caminho do plano de despesas excede 255 caracteres.")
                if subcaminho not in cache:
                    if pai and not CategoriaDespesa.objects.filter(caminho=subcaminho).exists():
                        pai = CategoriaDespesa.objects.select_for_update().get(pk=pai.pk)
                        verificar_pai_sem_vinculos(pai)
                    no, _ = CategoriaDespesa.objects.get_or_create(
                        caminho=subcaminho, defaults={"nome": nome, "nivel": nivel, "pai": pai},
                    )
                    if no.nome != nome or no.nivel != nivel or no.pai_id != (pai.pk if pai else None):
                        raise ValidationError(f"Árvore existente incompatível em {subcaminho!r}.")
                    cache[subcaminho] = no
                pai = cache[subcaminho]
            folhas.add(pai.pk)
            tipo, criado = TipoDespesa.objects.get_or_create(
                chave=chave(nome_tipo), defaults={"nome": nome_tipo, "codigo_origem": codigo},
            )
            if criado:
                auditar("tipo", tipo.pk, "CRIAR", depois={"nome": nome_tipo}, operador=operador)
            tipos.add(tipo.pk)
            salvar_vinculo(tipo, cache[caminho[0]], pai, operador)
    return {"folhas": len(folhas), "tipos": len(tipos), "linhas": len(itens)}


def numero_decimal(valor):
    if valor in (None, ""):
        return None
    try:
        numero = Decimal(str(valor).replace(",", ".")).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        return None
    return numero if numero.is_finite() else None


def data_planilha(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, str) and valor.strip():
        for formato in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(valor.strip(), formato).date()
            except ValueError:
                pass
    return None


def _texto_json(valor):
    if isinstance(valor, (date, datetime)):
        return valor.isoformat()
    if isinstance(valor, Decimal):
        return str(valor)
    return valor


def chave_base_dados(dados):
    partes = [dados.get("documento"), dados.get("emissao"), chave(dados.get("nome_tipo_origem")), chave(dados.get("fornecedor_texto"))]
    return sha256(json.dumps(partes, ensure_ascii=False).encode("utf-8")).hexdigest()


def hash_conteudo_dados(dados):
    conteudo = {k: dados.get(k) for k in (
        "nome_tipo_origem", "tipo_id", "fornecedor_texto", "documento", "observacoes",
        "emissao", "vencimento", "valor", "pagamento_data", "valor_pago", "forma",
    )}
    return sha256(json.dumps(conteudo, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()


def validar_linha(dados, numero, indices, tipos_por_chave, fornecedores_por_chave):
    def campo(*nomes):
        for nome in nomes:
            if nome in indices:
                return dados[indices[nome]] if indices[nome] < len(dados) else None
        return None

    identificador = str(campo("ID", "IDLANCAMENTO", "IDORIGEM", "IDDESPESA") or "").strip()
    nome_tipo = str(campo("DESPESA") or "").strip()
    tipo = tipos_por_chave.get(chave(nome_tipo)) if nome_tipo else None
    emissao_bruta = campo("EMISSAO")
    emissao = data_planilha(emissao_bruta)
    vencimento = data_planilha(campo("VENCIMENTO"))
    pagamento_bruto = campo("PAGTO", "PAGT", "PAGAMENTO", "DATAPAGAMENTO")
    pagamento = data_planilha(pagamento_bruto)
    valor = numero_decimal(campo("VALOR"))
    valor_pago_bruto = campo("VALORPAGO")
    valor_pago_convertido = numero_decimal(valor_pago_bruto)
    valor_pago = valor_pago_convertido if valor_pago_convertido is not None else Decimal("0")
    erros = []
    cabecalho_repetido = (
        chave(nome_tipo) == "DESPESA" and chave(str(campo("VENCIMENTO") or "")) == "VENCIMENTO"
    ) or (
        not nome_tipo and chave(emissao_bruta) == "DATA C" and chave(campo("VALOR")) == "VALOR TOTAL"
    )
    if cabecalho_repetido:
        erros.append("Cabeçalho repetido dentro dos dados")
    if not tipo:
        erros.append(f"Tipo de despesa não resolvido: {nome_tipo or '(vazio)'}")
    if emissao_bruta not in (None, "") and not emissao:
        erros.append("Data de emissão inválida")
    if not vencimento:
        erros.append("Vencimento inválido")
    if valor is None or valor < 0:
        erros.append("Valor inválido")
    if (valor_pago_bruto not in (None, "") and valor_pago_convertido is None) or valor_pago < 0 or (valor is not None and valor_pago > valor):
        erros.append("Valor pago inválido")
    if valor_pago > 0 and not pagamento:
        erros.append("Pagamento sem data válida")
    if pagamento_bruto not in (None, "") and not pagamento:
        erros.append("Data de pagamento inválida")
    fornecedor_texto = str(campo("FORNECEDOR") or "").strip()
    fornecedor = fornecedores_por_chave.get(chave(fornecedor_texto)) if fornecedor_texto else None
    normal = {
        "id_origem": identificador,
        "id_explicito": bool(identificador),
        "nome_tipo_origem": nome_tipo,
        "tipo_id": tipo.pk if tipo else None,
        "fornecedor_id": fornecedor.pk if fornecedor else None,
        "fornecedor_texto": fornecedor_texto,
        "documento": str(campo("DOC") or "").strip(),
        "observacoes": str(campo("OBSERVACOES") or "").strip(),
        "emissao": emissao.isoformat() if emissao else None,
        "vencimento": vencimento.isoformat() if vencimento else None,
        "valor": str(valor) if valor is not None else None,
        "pagamento_data": pagamento.isoformat() if pagamento else None,
        "valor_pago": str(valor_pago),
        "forma": str(campo("ESP") or "").strip(),
        "avisos": ["Fornecedor não vinculado; revise o cadastro se necessário"] if fornecedor_texto and not fornecedor else [],
    }
    normal["chave_base"] = chave_base_dados(normal)
    normal["hash_conteudo"] = hash_conteudo_dados(normal)
    return normal, tipo, erros


def atribuir_identidades(registros):
    """Concilia primeiro as ocorrências exatas; só atribui mudanças inequívocas."""
    existentes = defaultdict(list)
    por_id = {}
    hash_originais = Counter()
    for item in MovimentoDespesa.objects.filter(origem="XLSX").exclude(chave_base="").values(
        "id", "id_origem", "chave_base", "dados_origem", "documento", "vencimento", "valor", "tipo__nome",
    ).iterator(chunk_size=1000):
        item["hash_conteudo"] = (item["dados_origem"] or {}).get("hash_conteudo")
        existentes[item["chave_base"]].append(item)
        if item["id_origem"]:
            por_id[item["id_origem"]] = item
        hash_originais[(item["chave_base"], item["hash_conteudo"])] += 1
    for grupo in existentes.values():
        grupo.sort(key=lambda item: item["id"])

    usados_oficiais = set()
    usados_identidades = {}
    pendentes_por_base = defaultdict(list)

    def reservar(numero, dados, erros, oficial, situacao):
        identidade = oficial["id_origem"]
        if not identidade or identidade in usados_identidades:
            erros.append("Identidade ambígua: ocorrência oficial repetida no arquivo")
            dados["classificacao"] = "AMBIGUA"
            dados["identidade_situacao"] = "REVISAR"
            return
        dados["id_origem"] = identidade
        dados["movimento_id"] = oficial["id"]
        dados["classificacao"] = situacao
        dados["identidade_situacao"] = "CONTEUDO_EXATO" if situacao == "EXISTENTE" else "CONCILIADO_UNICO"
        usados_oficiais.add(oficial["id"])
        usados_identidades[identidade] = numero

    for numero, dados, _, erros in registros:
        if erros:
            dados["classificacao"] = "IGNORADA" if "Cabeçalho repetido dentro dos dados" in erros else "INVALIDA"
            continue
        if dados["id_explicito"]:
            oficial = por_id.get(dados["id_origem"])
            if oficial:
                situacao = "EXISTENTE" if oficial["hash_conteudo"] == dados["hash_conteudo"] else "ALTERADA"
                reservar(numero, dados, erros, oficial, situacao)
            else:
                dados["classificacao"] = "NOVA"
                dados["identidade_situacao"] = "ID_ORIGEM"
                if dados["id_origem"] in usados_identidades:
                    erros.append(f"Identidade repetida no arquivo (primeira linha {usados_identidades[dados['id_origem']]})")
                else:
                    usados_identidades[dados["id_origem"]] = numero
            continue
        pendentes_por_base[dados["chave_base"]].append((numero, dados, erros))

    # A identidade de uma repetição depende da quantidade de ocorrências, não da linha física.
    for base, pendentes in pendentes_por_base.items():
        por_hash = defaultdict(list)
        for oficial in existentes[base]:
            if oficial["id"] not in usados_oficiais:
                por_hash[oficial["hash_conteudo"]].append(oficial)
        sobras = []
        for numero, dados, erros in pendentes:
            correspondencias = por_hash[dados["hash_conteudo"]]
            if correspondencias:
                reservar(numero, dados, erros, correspondencias.pop(0), "EXISTENTE")
            else:
                sobras.append((numero, dados, erros))
        candidatos = [item for item in existentes[base] if item["id"] not in usados_oficiais]
        if len(candidatos) == len(sobras) == 1 and hash_originais[(base, candidatos[0]["hash_conteudo"])] == 1:
            numero, dados, erros = sobras[0]
            reservar(numero, dados, erros, candidatos[0], "ALTERADA")
        elif candidatos:
            opcoes = [{"movimento_id": c["id"], "tipo": c["tipo__nome"], "documento": c["documento"], "vencimento": c["vencimento"].isoformat(), "valor": str(c["valor"])} for c in candidatos[:100]]
            for _, dados, erros in sobras:
                dados["classificacao"] = "AMBIGUA"
                dados["identidade_situacao"] = "REVISAR"
                dados["candidatos"] = opcoes
                dados["total_candidatos"] = len(candidatos)
                erros.append("Identidade ambígua: revise antes de consolidar")
        else:
            for _, dados, _ in sobras:
                dados["classificacao"] = "NOVA"
                dados["identidade_situacao"] = "NOVO"

    # Gera IDs apenas para ocorrências ainda sem correspondência oficial.
    proximo = Counter()
    for numero, dados, _, erros in registros:
        if dados.get("classificacao") not in {"NOVA", "AMBIGUA", "INVALIDA"} or dados["id_explicito"]:
            continue
        prefixo = f"auto:{dados['hash_conteudo']}:"
        while True:
            proximo[prefixo] += 1
            identidade = f"{prefixo}{proximo[prefixo]}"
            if identidade not in por_id and identidade not in usados_identidades:
                break
        dados["id_origem"] = identidade
        usados_identidades[identidade] = numero
    for _, dados, _, erros in registros:
        if dados["classificacao"] not in {"AMBIGUA", "IGNORADA"} and erros:
            dados["classificacao"] = "INVALIDA"
    return registros


def criar_lote(arquivo, lote_existente=None):
    conteudo = arquivo.read()
    wb = load_workbook(BytesIO(conteudo), read_only=True, data_only=True)
    if "pgtdia" not in wb.sheetnames:
        raise ValidationError("A aba 'pgtdia' não foi encontrada. Corrija a planilha antes de importar.")
    linhas = wb["pgtdia"].iter_rows(min_row=4, values_only=True)
    try:
        nomes = next(linhas)
    except StopIteration as exc:
        raise ValidationError("A aba 'pgtdia' não possui cabeçalho na linha 4. Corrija a planilha antes de importar.") from exc
    indices = {cabecalho(nome): i for i, nome in enumerate(nomes)}
    if not {"DESPESA", "VALOR", "VENCIMENTO"}.issubset(indices):
        raise ValidationError("A linha 4 da aba 'pgtdia' deve conter DESPESA, VALOR e VENCIMENTO. Corrija a planilha antes de importar.")
    tipos = {tipo.chave: tipo for tipo in TipoDespesa.objects.all()}
    fornecedores = {}
    for fornecedor in Fornecedor.objects.all().only("id_fornecedor", "nome_fornecedor").iterator(chunk_size=1000):
        nome = chave(fornecedor.nome_fornecedor)
        if nome not in {"DESPESAS PLANO DE CONTAS", "FORNECEDORES DIVERSOS"}:
            if nome in fornecedores:
                fornecedores[nome] = None
            else:
                fornecedores[nome] = fornecedor
    arquivo_salvo = None
    try:
        with transaction.atomic():
            if lote_existente:
                lote = LoteDespesa.objects.select_for_update().get(pk=lote_existente.pk)
                if lote.status == "CONSOLIDADO" or lote.linhas.filter(status__in=["CONSOLIDADO", "CONFLITO"]).exists():
                    raise ValidationError("Um lote consolidado não pode ser revalidado.")
                lote.linhas.all().delete()
            else:
                if LoteDespesa.objects.exclude(status="CONSOLIDADO").exists():
                    raise ValidationError("Há uma captura pendente nesta filial. Conclua ou descarte antes de importar outra.")
                nome_arquivo = Path(arquivo.name).name[:255]
                lote = LoteDespesa(nome_arquivo=nome_arquivo, hash_sha256=sha256(conteudo).hexdigest(), chave_ativa="ATIVA")
                # A restrição única impede duas capturas concorrentes antes de salvar o arquivo.
                lote.save()
            registros = []
            for numero, linha in enumerate(linhas, 5):
                if not any(valor not in (None, "") for valor in linha):
                    continue
                normal, tipo, erros = validar_linha(linha, numero, indices, tipos, fornecedores)
                registros.append((numero, normal, tipo, erros))
            atribuir_identidades(registros)
            buffer = []
            for numero, normal, tipo, erros in registros:
                ignorada = "Cabeçalho repetido dentro dos dados" in erros
                buffer.append(LinhaDespesa(lote=lote, numero=numero, id_origem=normal["id_origem"], tipo=tipo, dados=normal, erros=erros, status="IGNORADO" if ignorada else "PENDENTE" if erros else "PRONTO"))
                if len(buffer) >= 1000:
                    LinhaDespesa.objects.bulk_create(buffer)
                    buffer.clear()
            if buffer:
                LinhaDespesa.objects.bulk_create(buffer)
            lote.total_linhas = LinhaDespesa.objects.filter(lote=lote).count()
            lote.prontas = LinhaDespesa.objects.filter(lote=lote, status="PRONTO").count()
            lote.pendentes = LinhaDespesa.objects.filter(lote=lote, status="PENDENTE").count()
            if lote_existente:
                lote.save(update_fields=["total_linhas", "prontas", "pendentes"])
            else:
                lote.arquivo.save(nome_arquivo, ContentFile(conteudo), save=False)
                arquivo_salvo = lote.arquivo.name
                lote.save(update_fields=["arquivo", "total_linhas", "prontas", "pendentes"])
    except Exception:
        if arquivo_salvo:
            LoteDespesa._meta.get_field("arquivo").storage.delete(arquivo_salvo)
        raise
    return lote


def meses_movimento(movimento):
    meses = {(mes_inicio(movimento.vencimento), "VENCIMENTO")}
    meses.update((mes_inicio(p.data), "PAGAMENTO") for p in movimento.pagamentos.all())
    return meses


def reconstruir_mes(mes, base):
    """Troca diário, mensal e snapshot juntos; falha preserva os agregados anteriores."""
    try:
        with transaction.atomic():
            snapshot, _ = SnapshotDespesaMensal.objects.select_for_update().get_or_create(mes=mes, base=base)
            fim = proximo_mes(mes)
            if base == "VENCIMENTO":
                valores = MovimentoDespesa.objects.filter(vencimento__gte=mes, vencimento__lt=fim).order_by().values("vencimento", "tipo_id").annotate(total=Sum("valor"))
            else:
                valores = PagamentoDespesa.objects.filter(data__gte=mes, data__lt=fim).order_by().values("data", "movimento__tipo_id").annotate(total=Sum("valor"))
            diario = []
            totais_tipo = defaultdict(Decimal)
            for linha in valores:
                tipo_id = linha["tipo_id"] if base == "VENCIMENTO" else linha["movimento__tipo_id"]
                dia = linha["vencimento"] if base == "VENCIMENTO" else linha["data"]
                diario.append(AgregadoDespesaDiario(dia=dia, tipo_id=tipo_id, base=base, valor=linha["total"]))
                totais_tipo[tipo_id] += linha["total"]
            mensal = [AgregadoDespesaMensal(mes=mes, tipo_id=tipo_id, base=base, valor=valor) for tipo_id, valor in totais_tipo.items()]
            AgregadoDespesaDiario.objects.filter(dia__gte=mes, dia__lt=fim, base=base).delete()
            AgregadoDespesaDiario.objects.bulk_create(diario, batch_size=500)
            AgregadoDespesaMensal.objects.filter(mes=mes, base=base).delete()
            AgregadoDespesaMensal.objects.bulk_create(mensal, batch_size=500)
            diario_gravado = dict(
                AgregadoDespesaDiario.objects.filter(dia__gte=mes, dia__lt=fim, base=base)
                .values("tipo_id").annotate(total=Sum("valor")).values_list("tipo_id", "total")
            )
            mensal_gravado = dict(AgregadoDespesaMensal.objects.filter(mes=mes, base=base).values_list("tipo_id", "valor"))
            if diario_gravado != mensal_gravado or mensal_gravado != dict(totais_tipo):
                raise ValidationError("Agregados diário e mensal de despesas não conciliam.")
            atualizado_em = timezone.now()
            snapshot.status = "PRONTO"
            snapshot.atualizado_em = atualizado_em
            snapshot.diario_atualizado_em = atualizado_em
            snapshot.erro = ""
            snapshot.save(update_fields=["status", "atualizado_em", "diario_atualizado_em", "erro"])
    except Exception as exc:
        SnapshotDespesaMensal.objects.update_or_create(
            mes=mes, base=base, defaults={"status": "FALHO", "erro": str(exc)[:1000]},
        )
        raise


def reconstruir_periodos(periodos):
    for mes, base in sorted(set(periodos)):
        reconstruir_mes(mes, base)


def consolidar_linha(linha: LinhaDespesa, operador="importacao"):
    if linha.erros:
        raise ValidationError("Linha possui pendências de validação.")
    d = linha.dados
    with transaction.atomic():
        mov = MovimentoDespesa.objects.select_for_update().filter(id_origem=d["id_origem"]).first()
        antes_periodos = meses_movimento(mov) if mov else set()
        campos = {"tipo_id": d["tipo_id"], "fornecedor_id": d["fornecedor_id"], "fornecedor_texto": d["fornecedor_texto"], "documento": d["documento"], "observacoes": d["observacoes"], "emissao": date.fromisoformat(d["emissao"]) if d["emissao"] else None, "vencimento": date.fromisoformat(d["vencimento"]), "valor": Decimal(d["valor"])}
        if mov is None:
            mov = MovimentoDespesa.objects.create(id_origem=d["id_origem"], chave_base=d.get("chave_base", ""), origem="XLSX", dados_origem=d, **campos)
            auditar("movimento", mov.pk, "IMPORTAR", depois=d, operador=operador)
        else:
            mudou = {}
            for campo, valor in campos.items():
                atual = getattr(mov, campo)
                if atual == valor:
                    continue
                if campo in mov.campos_manuais:
                    ConflitoDespesa.objects.update_or_create(
                        linha=linha, campo=campo,
                        defaults={"movimento": mov, "valor_atual": _texto_json(atual), "valor_origem": _texto_json(valor), "resolvido": False},
                    )
                    continue
                mudou[campo] = {"antes": _texto_json(atual), "depois": _texto_json(valor)}
                setattr(mov, campo, valor)
            if mudou:
                mov.save()
                auditar("movimento", mov.pk, "REIMPORTAR", antes={k: v["antes"] for k, v in mudou.items()}, depois={k: v["depois"] for k, v in mudou.items()}, operador=operador)
            mov.dados_origem = d
            mov.chave_base = d.get("chave_base", "")
            mov.save(update_fields=["dados_origem", "chave_base", "atualizado_em"])
        pagamento_id = f"{d['id_origem']}:1"
        pagamento = PagamentoDespesa.objects.select_for_update().filter(id_origem=pagamento_id).first()
        valor_pago = Decimal(d["valor_pago"])
        outros_pagamentos = sum((p.valor for p in mov.pagamentos.exclude(pk=pagamento.pk if pagamento else None)), Decimal("0"))
        if valor_pago > 0 and outros_pagamentos + (pagamento.valor if pagamento and pagamento.campos_manuais else valor_pago) > mov.valor:
            ConflitoDespesa.objects.update_or_create(
                linha=linha, campo="pagamento_total",
                defaults={"movimento": mov, "valor_atual": str(outros_pagamentos + (pagamento.valor if pagamento else Decimal("0"))), "valor_origem": str(valor_pago), "resolvido": False},
            )
            valor_pago = Decimal("0")
            pagamento = None
        if valor_pago > 0:
            novo = {"data": date.fromisoformat(d["pagamento_data"]), "valor": valor_pago, "forma": d["forma"]}
            if not pagamento:
                pagamento = PagamentoDespesa.objects.create(movimento=mov, id_origem=pagamento_id, origem="XLSX", **novo)
                auditar("pagamento", pagamento.pk, "IMPORTAR", depois={k: _texto_json(v) for k, v in novo.items()}, operador=operador)
            else:
                alterado = {}
                for campo, valor in novo.items():
                    atual = getattr(pagamento, campo)
                    if atual == valor:
                        continue
                    if campo in pagamento.campos_manuais:
                        ConflitoDespesa.objects.update_or_create(linha=linha, campo=f"pagamento_{campo}", defaults={"movimento": mov, "valor_atual": _texto_json(atual), "valor_origem": _texto_json(valor), "resolvido": False})
                    else:
                        setattr(pagamento, campo, valor)
                        alterado[campo] = {"antes": _texto_json(atual), "depois": _texto_json(valor)}
                if alterado:
                    pagamento.save()
                    auditar("pagamento", pagamento.pk, "REIMPORTAR", antes={k: v["antes"] for k, v in alterado.items()}, depois={k: v["depois"] for k, v in alterado.items()}, operador=operador)
        elif pagamento and not pagamento.campos_manuais:
            antes_periodos.add((mes_inicio(pagamento.data), "PAGAMENTO"))
            pagamento.delete()
        linha.movimento = mov
        linha.status = "CONFLITO" if ConflitoDespesa.objects.filter(linha=linha, resolvido=False).exists() else "CONSOLIDADO"
        linha.save(update_fields=["movimento", "status"])
        depois_periodos = meses_movimento(mov)
    return antes_periodos | depois_periodos


def consolidar_linhas_novas(linhas, operador):
    """Insere um bloco novo com poucas consultas, mantendo chaves idempotentes."""
    movimentos = []
    for linha in linhas:
        d = linha.dados
        movimentos.append(MovimentoDespesa(
            id_origem=d["id_origem"], chave_base=d.get("chave_base", ""),
            origem="XLSX", dados_origem=d, tipo_id=d["tipo_id"],
            fornecedor_id=d["fornecedor_id"], fornecedor_texto=d["fornecedor_texto"],
            documento=d["documento"], observacoes=d["observacoes"],
            emissao=date.fromisoformat(d["emissao"]) if d["emissao"] else None,
            vencimento=date.fromisoformat(d["vencimento"]), valor=Decimal(d["valor"]),
        ))
    MovimentoDespesa.objects.bulk_create(movimentos, batch_size=500)
    por_origem = MovimentoDespesa.objects.in_bulk([m.id_origem for m in movimentos], field_name="id_origem")
    pagamentos, auditorias, periodos = [], [], set()
    for linha in linhas:
        d = linha.dados
        mov = por_origem[d["id_origem"]]
        linha.movimento_id = mov.pk
        linha.status = "CONSOLIDADO"
        periodos.add((mes_inicio(mov.vencimento), "VENCIMENTO"))
        auditorias.append(AuditoriaDespesa(entidade="movimento", entidade_id=mov.pk, acao="IMPORTAR", depois=d, operador=operador))
        valor_pago = Decimal(d["valor_pago"])
        if valor_pago > 0:
            data_pagamento = date.fromisoformat(d["pagamento_data"])
            pagamentos.append(PagamentoDespesa(
                movimento_id=mov.pk, id_origem=f"{d['id_origem']}:1", origem="XLSX",
                data=data_pagamento, valor=valor_pago, forma=d["forma"],
            ))
            periodos.add((mes_inicio(data_pagamento), "PAGAMENTO"))
    PagamentoDespesa.objects.bulk_create(pagamentos, batch_size=500)
    por_pagamento = PagamentoDespesa.objects.in_bulk([p.id_origem for p in pagamentos], field_name="id_origem")
    for pagamento in pagamentos:
        gravado = por_pagamento[pagamento.id_origem]
        auditorias.append(AuditoriaDespesa(
            entidade="pagamento", entidade_id=gravado.pk, acao="IMPORTAR",
            depois={"data": gravado.data.isoformat(), "valor": str(gravado.valor), "forma": gravado.forma}, operador=operador,
        ))
    AuditoriaDespesa.objects.bulk_create(auditorias, batch_size=500)
    LinhaDespesa.objects.bulk_update(linhas, ["movimento", "status"], batch_size=500)
    return periodos


def consolidar_lote(lote, operador="importacao"):
    periodos = set()
    bloqueio_adquirido = False
    try:
        with transaction.atomic():
            try:
                lote = LoteDespesa.objects.select_for_update(nowait=connection.features.has_select_for_update_nowait).get(pk=lote.pk)
            except DatabaseError as exc:
                raise ValidationError("Este lote já está sendo consolidado. Aguarde a conclusão e atualize a tela.") from exc
            bloqueio_adquirido = True
            if lote.status == "CONSOLIDADO":
                return {"linhas": lote.total_linhas, "conflitos": lote.linhas.filter(status="CONFLITO").count(), "periodos": 0, "ja_consolidado": True}
            if lote.pendentes:
                raise ValidationError("Resolva as linhas pendentes antes de consolidar o lote.")
            ultima_linha = lote.linhas.filter(status="PRONTO").aggregate(numero=Max("numero"))["numero"] or 0
            for inicio in range(0, ultima_linha, 500):
                linhas = list(lote.linhas.filter(status="PRONTO", numero__gte=inicio + 1, numero__lt=inicio + 501).order_by("numero"))
                if not linhas:
                    continue
                ids = [linha.id_origem for linha in linhas]
                existentes = MovimentoDespesa.objects.in_bulk(ids, field_name="id_origem")
                manuais_pagamento = set(PagamentoDespesa.objects.filter(id_origem__in=[f"{identidade}:1" for identidade in ids]).exclude(campos_manuais=[]).values_list("id_origem", flat=True))
                novas, iguais, alteradas = [], [], []
                for linha in linhas:
                    mov = existentes.get(linha.id_origem)
                    if mov is None:
                        novas.append(linha)
                    elif (mov.dados_origem.get("hash_conteudo") == linha.dados.get("hash_conteudo")
                          and not mov.campos_manuais and f"{linha.id_origem}:1" not in manuais_pagamento):
                        linha.movimento_id = mov.pk
                        linha.status = "CONSOLIDADO"
                        iguais.append(linha)
                    else:
                        alteradas.append(linha)
                if novas:
                    periodos |= consolidar_linhas_novas(novas, operador)
                if iguais:
                    LinhaDespesa.objects.bulk_update(iguais, ["movimento", "status"], batch_size=500)
                for linha in alteradas:
                    periodos |= consolidar_linha(linha, operador)
            reconstruir_periodos(periodos)
            lote.status = "CONSOLIDADO"
            lote.chave_ativa = None
            lote.save(update_fields=["status", "chave_ativa"])
    except Exception:
        if bloqueio_adquirido:
            LoteDespesa.objects.filter(pk=lote.pk).update(status="FALHO")
        raise
    return {"linhas": lote.total_linhas, "conflitos": lote.linhas.filter(status="CONFLITO").count(), "periodos": len(periodos)}
